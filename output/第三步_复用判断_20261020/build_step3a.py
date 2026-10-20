"""③a 全覆盖判断：17 类方法逐类判定能否复用到固态锂硫（摘要级）。
判定文字是 Claude 依据摘要原句写的，用户未复核；凡标"推理"的是没有原文直接支持的判断。
证据按（论文号或 DOI，关键短语）给出，代码在规范化的标题/摘要里找到含该短语的整句并回查。
"""
import json, re, sys, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
INV = OUT.parent / '方法清单_硫转化动力学_20261019'
sys.path.insert(0, str(INV))
from common import norm, load_corpus, sentences
from build_inventory import VOCAB

sys.stdout.reconfigure(encoding='utf-8')
ev = json.load(open(OUT / '_step3a_evidence.json', encoding='utf-8'))
inv = json.load(open(INV / '_inventory.json', encoding='utf-8'))
corpus = load_corpus()
by_id = {r['id']: r for r in corpus}
by_doi = {r['doi'].lower(): r for r in corpus if r.get('doi')}

J = {
'M01': dict(
  解决='液态：原子级金属位点吸附多硫化物、降低转化能垒；摘要 88% 讲吸附/锚定/穿梭，40% 讲固相步骤（Li₂S 成核/分解等）',
  前提='① 硫物种能到达位点（液态靠溶解的多硫化物扩散）；② 位点与硫紧密接触且有电子供给（负载在导电碳上）',
  硫化物='① 不成立（没有溶解的多硫化物）；② 可成立，但只在硫｜碳｜电解质接触处', 聚合物='① 部分成立（PEO 中多硫化物可溶）',
  已有='全部 2023 年后：M–S 键强度作为双向催化指标、Ni 单原子调控 Li₂S₂ 还原、Co 单原子串联催化、三相界面定向催化、双原子电子缓冲',
  判定='需改造（固态锂硫已在快速推进）',
  改造='筛选判据从"吸附溶解的多硫化物"换成"与固态硫物种的结合强度/作用于固相步骤"；位点须在三相界面；须核对是否加速硫化物电解质分解',
  缺='硫化物与聚合物体系分开的效果；与电解质分解的定量权衡；同条件下不同单原子的比较', 建议='高（与 M02 合并为"催化剂"）',
  证据=[('W4402207230', 'bond formed between sacs and sulfur', '固态已有'), ('W4407383031', 'nickel single-atom catalyst', '固态已有'),
        ('W7127084998', 'triple phase interface', '固态已有'), ('AUTO:液态锂硫:M01', 'adsor', '原体系逻辑')]),
'M02': dict(
  解决='液态：金属化合物/异质结构/缺陷位点吸附并催化多硫化物转化；摘要 88% 讲吸附/锚定/穿梭，26% 讲固相步骤',
  前提='同 M01；另需催化剂与硫化物电解质化学相容',
  硫化物='吸附锚定不成立；降低固相步骤能垒可成立；已有报道催化剂会加速硫化物电解质分解、过渡金属迁移致电解质降解', 聚合物='部分成立',
  已有='CoS₂ 均质复合正极、TiS₂ 催化"焊料"连接界面、Co 尺寸与电解质分解的权衡、高熵硫化物催化剂（指出自中毒和过渡金属迁移两种失效）',
  判定='需改造（固态锂硫已在推进）',
  改造='同 M01，并把"不加速电解质分解"列为筛选条件（尺寸、空间隔离、相稳定）',
  缺='同 M01；各材料族在硫化物体系中的相容性', 建议='高（与 M01 合并）',
  证据=[('W7169438107', 'physically isolating the catalytic surfaces', '风险与改造'), ('W7203675724', 'electrolyte degradation', '风险'),
        ('W4411909838', 'interfacial continuity', '固态已有'), ('AUTO:液态锂硫:M02', 'adsor', '原体系逻辑')]),
'M03': dict(
  解决='液态：可溶介体在电极与绝缘的 S/Li₂S 之间传递电子',
  前提='介体能在电解质中扩散',
  硫化物='不成立；固态论文已改为界面介体、固定化介体和氧化还原活性 catholyte', 聚合物='部分成立',
  已有='2021 年起，主要用于活化 Li₂S 正极、降低首次充电电压（LiVS₂ 壳、有机硒盐、二苯基硫族化合物、硫代锡酸锂）；有工作提醒电解质作介体时会严重降解',
  判定='需改造（已改造，以 Li₂S 正极为主）', 改造='介体只能在接触点起作用，需固定在界面或成为电解质组分',
  缺='对 S₈ 正极放电动力学的作用证据少；介体与硫化物电解质的长期稳定性', 建议='中',
  证据=[('10.1002/smtd.202501077', 'reduce the initial charge voltage', '固态已有'), ('10.1039/d2ee03297j', 'redox mediator', '固态已有'),
        ('10.1021/acs.chemmater.2c02926', 'severe degradation of electrolytes', '风险'), ('AUTO:液态锂硫:M03', 'soluble', '原体系逻辑')]),
'M04': dict(
  解决='液态：SPAN、共价硫、小分子硫不经溶解的多硫化物，走固-固转化；Li₂S 正极',
  前提='反应不依赖溶解的多硫化物', 硫化物='成立', 聚合物='成立',
  已有='固态锂硫自己最常用的路线：Li₂S 正极（纳米化、Li₂S–LiI 固溶体、原位 Li₃PS₄ 包覆）、SPAN（多配聚合物电解质）、Te/Se 掺杂、多硫化物正极',
  判定='可直接复用（前提成立）', 改造='液态里针对 SPAN 动力学的改进（Se 掺杂、在 SPAN 上引入催化位点）可直接借；SPAN 在硫化物全固态中研究较少',
  缺='SPAN 在硫化物全固态中的离子传输与界面相容性；硫含量低带来的能量密度代价', 建议='中',
  证据=[('10.1021/acsaem.2c00978', 'ionic conductive domains', '固态已有'), ('10.1039/d1se01187a', 'sulfurized polyacrylonitrile', '固态已有'),
        ('10.1002/aesr.70223', 'covalently bound sulfur', '原体系做法')]),
'M05': dict(
  解决='液态：调溶剂化结构和添加剂，改变溶解物种的转化路径', 前提='液相溶剂化',
  硫化物='不成立', 聚合物='部分成立（Li⁺ 与聚合物链配位，归入 M11 讨论）',
  已有='3 篇（叠氮化锂、CsGeI₃ 作电解质添加剂等）', 判定='不适用（硫化物全固态）', 改造='—', 缺='—', 建议='低',
  证据=[('10.1002/anie.201709305', 'electrolyte additive', '固态已有'), ('AUTO:液态锂硫:M05', 'solvation', '原体系逻辑')]),
'M06': dict(
  解决='液态：多孔碳宿主限域多硫化物并提供电子通路；摘要 91% 讲吸附/锚定/穿梭',
  前提='① 液态电解液能浸入孔道；② 碳提供电子通路',
  硫化物='① 不成立（固态电解质进不了孔，已改为把硫和电解质溶液一起渗入介孔碳）；② 成立', 聚合物='部分成立',
  已有='2013 年起广泛采用；固态特有发现：碳维度比较、团聚敏感性可作设计指标、碳比表面积越大电解质分解越多',
  判定='需改造（已广泛采用）', 改造='孔结构要服务于"电解质能到达"而不是"限域"；权衡比表面积与电解质分解',
  缺='同条件下碳类型、比表面、孔结构的规律', 建议='中',
  证据=[('10.1016/j.xcrp.2024.102293', 'aggregation sensitivity', '固态已有'), ('10.1039/d3ya00513e', 'solid electrolyte solutions into the interior of the porous carbon', '改造做法'),
        ('10.1021/acsaem.5c02038', 'distinct dimensionalities', '固态已有'), ('W7124424515', 'carbon host nanostructure engineering', '风险')]),
'M07': dict(
  解决='液态：导电聚合物包覆或修饰宿主，导电并阻挡多硫化物',
  前提='在硫正极工作电位下保持导电（掺杂态）；与硫化物电解质相容',
  硫化物='不确定。推理：PPy、PEDOT 等在低电位可能脱掺杂而失去导电性（无原文证据）', 聚合物='不确定',
  已有='2 篇（2014；2024 弹性导电通道）；固态锂金属里 PEDOT:PSSTFSI 作 NCM 正极的混合导体 catholyte 粘结剂',
  判定='不确定（关键前提未核）', 改造='—', 缺='导电聚合物在约 1.5–2.8 V 下的导电性与稳定性；与硫化物电解质的相容性', 建议='中（篇数少是空白线索，但前提风险大）',
  证据=[('10.1002/aenm.202304412', 'conducting polymer', '固态已有'), ('10.1002/aenm.202506757', 'pedot', '固态锂金属做法')]),
'M08': dict(
  解决='固态锂金属多用于负极中间层，少数作 NCM 正极 catholyte（PEDOT:PSSTFSI + Li₃InCl₆）；液态用作中间层',
  前提='混合导体在硫正极电位下稳定、与硫物种相容；作用是把三相界面变成两相界面',
  硫化物='不确定。推理：含 Ti⁴⁺ 的 LLTO、LATP 在低电位可能被还原', 聚合物='不确定',
  已有='4 篇：混合导电 Li₂S 纳米复合（2016）、LLTO/C 两相界面（2021）、C/LATP 壳（2024）、离子液体锚定 CNT 网络（2025）',
  判定='不确定（方向同时对准离子和电子两条原因）', 改造='—', 缺='在硫正极电位下的稳定性；与 catholyte、碳的分工', 建议='高',
  证据=[('10.1021/acs.nanolett.1c04228', 'double-phase interface', '固态已有'), ('10.1002/aenm.202302255', 'mixed conductive', '固态已有'),
        ('10.1021/acsenergylett.5c03148', 'mixed ion-electron conductor network', '固态已有')]),
'M09': dict(
  解决='固态锂金属：正极电解质与隔膜电解质分开选（卤化物 catholyte、正负极用不同硫银锗矿）、厚 NCM 正极里的三维离子网络、混合导体 catholyte',
  前提='为 NCM 选 catholyte 的判据是抗氧化；硫正极需要抗还原，并与 S/Li₂S/多硫化物相容',
  硫化物='部分成立（判据要换）：Li₃InCl₆ 抗氧化好，却在硫正极中因还原快速降解；Li₃YBr₆ 表现最好', 聚合物='交联聚合物 catholyte 有 1 篇',
  已有='2023 年后为主：Li₃YCl₅I 卤化物 catholyte、氧化还原活性 catholyte、W 取代硫化物 catholyte、catholyte 部分离子传输的作用、梯度 catholyte 浓度',
  判定='需改造（选材判据从抗氧化换成抗还原 + 与硫物种相容）',
  改造='主因证据：中子成像两篇显示复合正极内离子传输是限速步，是目前直接证据最多的原因',
  缺='候选 catholyte 在硫正极电位下的还原稳定性数据；高载量下离子通路的分布', 建议='高',
  证据=[('10.1002/adma.202511553', 'critical factors limiting sulfur conversion', '主因证据'), ('10.1002/aenm.202203426', 'rate-limiting', '主因证据'),
        ('W4405443748', 'rapid degradation due to electrolyte reduction', '前提核对'), ('10.1021/acs.nanolett.4c05892', 'catholyte', '固态已有'),
        ('10.1002/smtd.202501793', 'catholyte', '固态锂金属做法'), ('10.1021/acsaem.5c04044', 'ionic transport network', '固态锂金属做法')]),
'M10': dict(
  解决='固态锂金属：掺杂、取代提高电导和对锂负极的稳定性', 前提='材料体系相同；对正极动力学的作用只在用作 catholyte 时直接',
  硫化物='成立', 聚合物='—', 已有='掺杂 Li₇P₃S₁₁、硫银锗矿、高熵硫银锗矿；In₂S₃ 掺杂抑制电化学分解',
  判定='可直接复用（作为 catholyte 候选的来源，并入 M09）', 改造='—', 缺='改性电解质用在正极内的稳定性', 建议='低（并入 M09）',
  证据=[('10.1021/acsaem.2c02110', 'suppresses electrochemical decomposition', '固态已有'), ('10.1016/j.xcrp.2024.102228', 'high-entropy', '固态已有')]),
'M11': dict(
  解决='固态锂金属：提高离子电导、稳定锂负极', 前提='材料相同；但聚合物锂硫中多硫化物会溶解，常需约 60 °C 运行，反应不是固-固转化',
  硫化物='—', 聚合物='成立', 已有='多为 PEO 体系',
  判定='可直接复用（作电解质层），与主线关系弱', 改造='—', 缺='—', 建议='低',
  证据=[('10.1021/acs.macromol.5b00928', 'lithium polysulfides', '前提核对'), ('10.1002/adfm.202203272', 'polysulfide fixation', '固态已有')]),
'M12': dict(
  解决='固态锂金属：离子液体、增塑剂、原位聚合提高电导和界面润湿', 前提='体系含液相成分',
  硫化物='与全固态定义冲突', 聚合物='成立，但成为准固态', 已有='多为聚合物体系；凝胶包封硫正极等',
  判定='不适用于全固态主线（绕开了固-固转化，可作对照）', 改造='—', 缺='—', 建议='低',
  证据=[('10.1149/1945-7111/ac64cb', 'gel polymer electrolyte', '固态已有'), ('10.5229/jkes.2010.13.2.110', 'plasticized', '固态已有')]),
'M13': dict(
  解决='各体系通用的合成与混合工艺', 前提='通用', 硫化物='成立', 聚合物='成立',
  已有='球磨减小粒径以减轻体积变化、两步混合、液相合成、溶液渗入',
  判定='可直接复用（固态锂硫已成熟，跨体系增量小；作为任何方向的实施手段）', 改造='—', 缺='—', 建议='低（工具）',
  证据=[('10.1021/acs.chemmater.9b00282', 'reducing the particle size is key', '固态已有'), ('10.1002/aenm.202500867', 'two-step mixing', '固态已有')]),
'M14': dict(
  解决='液态：自支撑、三维、低曲折度孔道（靠电解液灌满）；固态锂金属：厚 NCM 正极的三维离子网络、电解质双层/夹层',
  前提='液态孔道由电解液填充；固态锂金属的三维离子网络由固态电解质构成',
  硫化物='液态做法不成立；固态锂金属做法成立', 聚合物='部分成立',
  已有='梯度正极（与中子成像同一工作）、浓度梯度硫正极、双界面正极、干法自支撑正极片',
  判定='需改造（并入 M09 讨论"正极内离子通路"）', 改造='—', 缺='高载量下结构对离子通路的定量影响', 建议='中—高（与 M09 合并）',
  证据=[('10.1002/adma.202511553', 'gradient cathode architectures', '固态已有'), ('10.1002/adfm.202530425', 'concentration-gradient', '固态已有'),
        ('10.1002/cctc.70811', 'tortuosity', '原体系做法')]),
'M15': dict(
  解决='液态：隔膜涂层/中间层挡多硫化物、纳米封装；固态锂金属：正极颗粒包覆，稳定正极｜电解质界面',
  前提='液态挡穿梭靠溶解的多硫化物；NCM 包覆以抗氧化为目标',
  硫化物='挡穿梭不成立；包覆需改为降低界面离子转移阻力、抑制电解质分解', 聚合物='封装概念已被迁移（2020）',
  已有='Li₂S@Li₃PS₄ 核壳、LiVS₂ 壳、ALD 涂层、硫正极纳米尺度界面工程、机械诱导桥接中间层',
  判定='需改造（目标换成降低硫/Li₂S｜电解质界面的离子转移阻力）',
  改造='主因证据：Li₆PS₅Cl｜Li₂S 界面的离子传输比电解质体相小几个数量级',
  缺='界面层对界面离子转移的定量改善；与电解质分解的关系', 建议='中—高',
  证据=[('10.1021/jacs.6b05066', 'orders of magnitude smaller', '主因证据'), ('10.1021/acs.nanolett.0c02033', 'encapsulation concept from liquid electrolytes', '跨体系迁移实例'),
        ('10.1002/adfm.202315925', 'interface engineering of sulfur cathode', '固态已有'), ('AUTO:液态锂硫:M15', 'separator', '原体系做法')]),
'M16': dict(
  解决='液态：功能粘结剂；固态锂金属：正极里的锂离子导电粘结剂、功能聚合物粘结剂',
  前提='硫与 Li₂S 之间的体积变化导致接触丢失；软质或导离子粘结剂能维持接触，这一点在固态更关键',
  硫化物='成立（需核粘结剂的还原稳定性）', 聚合物='成立',
  已有='柔顺正极、弹性导电通道、无机离子导电粘结剂 nano-Li₃PS₄、弹性体电解质、低压下粘结剂对接触的影响',
  判定='需改造（从 NCM 正极迁移导离子粘结剂，核还原稳定性）', 改造='主因证据：有工作指出体积变化不可逆是硫正极容量衰减的根源',
  缺='接触保持对转化动力学的定量影响（区别于循环衰减）', 建议='中',
  证据=[('10.1021/acs.nanolett.6c00075', 'origin of the capacity fading', '主因证据'), ('10.1021/acsami.3c16731', 'ion-conductive binder', '固态已有'),
        ('10.1002/adma.202520337', 'li-conductive binder', '固态锂金属做法')]),
'M17': dict(
  解决='—（工况）', 前提='—', 硫化物='—', 聚合物='—', 已有='低压、室温运行的演示',
  判定='工况条件，不判复用；作为检验条件（低压、室温）', 改造='—', 缺='—', 建议='不进入',
  证据=[('10.1002/adfm.202423633', 'low pressure and ambient temperature', '固态已有'), ('10.1021/acs.chemmater.4c01514', 'stacking pressure', '固态已有')]),
}

DIAG = [('10.1002/aenm.202203426', 'rate-limiting', '离子', '中子成像'), ('10.1002/adma.202511553', 'critical factors limiting sulfur conversion', '离子', '原位中子成像'),
        ('10.1002/aenm.202503863', 'transport bottleneck', '离子', '高分辨可视化'), ('10.1021/jacs.6b05066', 'orders of magnitude smaller', '离子（界面）', '界面离子传输测量'),
        ('10.1002/anie.202302363', 'remained at the full lithiation stage', '反应（作者归因）', '原位与理论'),
        ('10.1002/aenm.202405642', 'incomplete conversion of li2s2 to li2s', '反应（作者归因）', '—'),
        ('10.1021/acs.nanolett.6c00075', 'origin of the capacity fading', '接触/体积', '原位研究（针对衰减）')]


def find(key, phrase):
    if key.startswith('AUTO:'):
        _, scope, code = key.split(':')
        for r in inv['rows']:
            if r['范围'] == scope and r['类别编号'] == code and phrase in r['原句（规范化）']:
                return by_id[r['编号']], r['原句（规范化）']
        return None, None
    rec = by_id.get(key) or by_doi.get(key.lower())
    if not rec:
        return None, None
    for s in [norm(rec['title'])] + sentences(rec['abstract']):
        if phrase in s:
            return rec, s
    return rec, None


ev_rows, n_ok = [], 0
for code, j in list(J.items()):
    for key, ph, use in j['证据']:
        rec, s = find(key, ph)
        ok = bool(rec and s)
        n_ok += ok
        ev_rows.append([code, use, rec['id'] if rec else key, rec['year'] if rec else '', rec['doi'] if rec else '',
                        rec['group'] if rec else '', ph, s or '（未找到）', '✅' if ok else '❌'])
diag_rows = []
for key, ph, cause, how in DIAG:
    rec, s = find(key, ph)
    diag_rows.append([cause, how, rec['id'] if rec else key, rec['year'] if rec else '', rec['doi'] if rec else '', s or '（未找到）', '✅' if (rec and s) else '❌'])
print('证据回查', n_ok, '/', len(ev_rows), '；主因证据', sum(r[-1] == '✅' for r in diag_rows), '/', len(diag_rows))
for r in ev_rows + diag_rows:
    if r[-1] == '❌': print('  未找到：', r[:3], r[-3] if len(r) > 7 else '')

# ---------------- 复核表 ----------------
wb = Workbook()
ws = wb.active
ws.title = '说明'
for i, (t, b) in enumerate([
    ('第 ③a 步：17 类方法能否复用到固态锂硫（摘要级判断）', True),
    ('主问题：固态锂硫正极的硫转化动力学受限。本表覆盖四条原因下的全部 17 类方法，只用摘要和方法清单原句，未读全文、未下载。', False),
    ('判定为 Claude 依据原句写出，用户未复核。标"推理"的内容没有原文直接支持，需在 ③b 核实。', False),
    ('判定只有四种：可直接复用（前提在固态硫正极中成立）/ 需改造（写清改什么）/ 不适用 / 不确定（关键前提未核）。', False),
    ('"③b 建议"是 Claude 的建议，最终由用户挑 2–4 类进入 ③b 深读，④ 只做其中 1 类。', False),
    ('', False),
    ('页签', True),
    ('③a判定表：每类一行。主因证据：四条原因里哪条是主因的诊断原句。证据原句：判定所引原句，代码逐字回查（✅ 找到 / ❌ 未找到）。篇数与前提线索：三体系篇数和"前提线索"比例（摘要提到某类词的比例，只是线索）。', False),
    ('', False),
    (f'证据回查：{n_ok}/{len(ev_rows)} 条找到；主因证据 {sum(r[-1] == "✅" for r in diag_rows)}/{len(diag_rows)} 条找到。', False),
], 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(bold=b, size=12 if b else 11)
    c.alignment = Alignment(wrap_text=True)
ws.column_dimensions['A'].width = 130

ws = wb.create_sheet('③a判定表')
head = ['编号', '方法类别', '针对的原因', '固态锂硫篇数（2023后）', '液态锂硫篇数', '固态锂金属篇数', '原体系解决什么', '起效前提',
        '前提在固态硫正极·硫化物体系', '前提在固态硫正极·聚合物体系', '固态锂硫已有工作', '判定', '改造点 / 主因证据', '还缺什么证据', '③b 建议', '人工确认✍']
ws.append(head)
FILL = {'可直接复用': 'DCEFD9', '需改造': 'FFF1CC', '不适用': 'EEEEEE', '不确定': 'F9DCDC', '工况': 'EEEEEE'}
for code, cause, name, _ in VOCAB:
    if code == 'X01': continue
    j, e = J[code], ev[code]
    row = [code, name, cause, f"{e['固态锂硫']['篇数']}（{e['固态锂硫']['2023后']}）", e['液态锂硫']['篇数'], e['固态锂金属']['篇数'],
           j['解决'], j['前提'], j['硫化物'], j['聚合物'], j['已有'], j['判定'], j['改造'], j['缺'], j['建议'], '']
    ws.append(row)
    for k, col in FILL.items():
        if j['判定'].startswith(k):
            ws.cell(row=ws.max_row, column=12).fill = PatternFill('solid', fgColor=col)
for i, w in enumerate([6, 26, 12, 11, 9, 9, 38, 34, 34, 22, 40, 22, 40, 34, 16, 10], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DCE9F7'); c.alignment = Alignment(wrap_text=True, vertical='center')
for r in ws.iter_rows(min_row=2):
    for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'

ws = wb.create_sheet('主因证据')
ws.append(['指向的原因', '证据类型', '论文号', '年份', 'DOI', '原句（规范化）', '回查'])
for r in diag_rows: ws.append(r)
ws.append([])
ws.append(['小结', '', '', '', '', '有直接实验证据（成像/界面测量）的 4 条都指向离子传输（复合正极宏观传输、硫/Li₂S｜电解质界面）；指向反应本身的 2 条是作者归因，未排除传输影响。', ''])
for i, w in enumerate([14, 18, 13, 6, 28, 110, 6], 1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)

ws = wb.create_sheet('证据原句')
ws.append(['类别', '用途', '论文号', '年份', 'DOI', '所属组', '关键短语', '原句（规范化）', '回查'])
for r in ev_rows: ws.append(r)
for i, w in enumerate([6, 12, 13, 6, 28, 10, 26, 110, 6], 1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)
ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet('篇数与前提线索')
ws.append(['编号', '方法类别', '固态锂硫篇数', '固态锂硫2023后', '固态锂硫最早年份', '固态锂硫电解质类型', '液态锂硫篇数', '液态·前提线索', '固态锂金属篇数', '固态锂金属·前提线索'])
for code, cause, name, _ in VOCAB:
    if code == 'X01': continue
    e = ev[code]
    fmt = lambda dct: '；'.join(f'{k.split("·")[1]} {v:.0%}' for k, v in dct.items() if v is not None)
    ws.append([code, name, e['固态锂硫']['篇数'], e['固态锂硫']['2023后'], e['固态锂硫']['最早年份'],
               '；'.join(f'{k} {v}' for k, v in e['固态锂硫']['电解质类型'].items()),
               e['液态锂硫']['篇数'], fmt(e['液态锂硫']['前提线索（论文占比）']), e['固态锂金属']['篇数'], fmt(e['固态锂金属']['前提线索（论文占比）'])])
for i, w in enumerate([6, 30, 9, 9, 9, 40, 9, 60, 9, 70], 1): ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]: c.font = Font(bold=True)
wb.save(OUT / '第三步a_复用判定表.xlsx')
print('写出 第三步a_复用判定表.xlsx')
print('判定分布', collections.Counter(J[c]['判定'].split('（')[0] for c in J))
