"""生成 F 定向核查的复核表：逐篇判定 + 原句核对。

输入：F_逐篇判定_待核.csv（人工判定）、F_筛选分层.csv（f_screen.py 的分层结果）、
      _全文清单.csv（论文号 → PDF 路径）、冻结论文表。
做的事：T1/T2 里未列入判定表的论文补成"不计"行并写明原因；为有原句的论文抽带页码的全文文本；
        逐字核对原句（忽略空白、连字符、连字和大小写差异）；写出 F_逐篇复核表.xlsx 和 F_汇总.json。

用法：python build_f_table.py <F_定向核查目录> <dataset_v2 目录>
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCLUDED = {
    'W7117979440': '负极侧', 'W7212020979': '边界：金属多硫化物正极', 'W7213290624': '无关：电镜综述',
    'W4410547199': '无关：SECCM 综述', 'W4401308475': '负极侧', 'W3080132686': '负极侧',
    'W2913315989': '无关：液固通用展望', 'W2895629147': '负极侧',
    'W7140296670': '不计：主题是多硫化物在电解质中的扩散', 'W7156377761': '无关：含锌硫体系',
    'W7171946273': '无关：导电聚合物综述', 'W7214333712': '负极侧', 'W4407675385': '负极侧',
    'W4409139349': '边界：PEO 与 LPSCl 的化学分解，不是正极中的氧化还原',
    'W4409810847': '边界：FeS2 正极，按口径属硫系转化型', 'W4415098663': '无关：以液态、准固态为主',
    'W4393390359': '负极侧', 'W4401629105': '边界：钠体系', 'W4403954692': '无关：锂金属电池综述',
    'W4405729176': '顺带提及', 'W4387002533': '负极侧', 'W4223953941': '负极侧', 'W4285283930': '负极侧',
    'W4309047134': '无关：聚合物综述未涉及', 'W4312043050': '无关：液态', 'W3131880064': '无关：Li–O2 等',
    'W3134932327': '无关：硅电极综述', 'W3178399305': '无关：液态锂金属', 'W3002270725': '顺带提及',
    'W2908453675': '边界：VS4 正极', 'W2993271346': '顺带提及', 'W2996752519': '无关：液态',
    'W2840670318': '无关：液态电解液综述', 'W2904499258': '无关：贡献容量的是 Se，不是电解质',
    'W2029872662': '无关',
}
FILL = {'抑制': 'DCE7F7', '权衡': 'FCE4D6', '利用': 'D7F0E6', '外加介体': 'FFF2CC',
        '量化或校正': 'F2E3EC', '只观察': 'F2E3EC', '综述': 'EEEEEE', '不计': 'FFFFFF'}


def norm(s):
    return re.sub(r'[\s\-‐‑‒–—―−­]', '', unicodedata.normalize('NFKC', s)).lower()


def page_text(pdf, out):
    if not out.exists():
        d = pymupdf.open(pdf)
        out.write_text(''.join(f'=== 第 {i} 页 ===\n{p.get_text()}\n' for i, p in enumerate(d, 1)), encoding='utf-8')
        d.close()
    buf, spans, pos, cur = [], [], 0, None
    for part in re.split(r'(=== 第 \d+ 页 ===)', out.read_text(encoding='utf-8')):
        m = re.fullmatch(r'=== 第 (\d+) 页 ===', part)
        if m:
            cur = int(m.group(1)); continue
        s = norm(part); buf.append(s); spans.append((pos, pos + len(s), cur)); pos += len(s)
    return ''.join(buf), spans


def main(d, ds):
    d, ds = Path(d), Path(ds)
    jd = pd.read_csv(d / 'F_逐篇判定_待核.csv', dtype=str).fillna('')
    sc = pd.read_csv(d / 'F_筛选分层.csv', dtype=str).fillna('').set_index('编号')
    mf = pd.read_csv(d / '_全文清单.csv', dtype=str).set_index('编号')
    pp = pd.read_csv(ds / '论文表.csv', dtype=str, low_memory=False).fillna('').set_index('编号')
    read = sc[sc['tier'].isin(['T1', 'T2'])].index
    miss = [w for w in read if w not in set(jd['论文号'])]
    assert set(miss) == set(EXCLUDED), (set(miss) ^ set(EXCLUDED))
    ex = pd.DataFrame([{'论文号': w, '立场': '不计', '备注': EXCLUDED[w]} for w in miss])
    df = pd.concat([jd, ex], ignore_index=True).fillna('')
    df.insert(1, '年份', df['论文号'].map(lambda w: pp.at[w, '年份']))
    df.insert(2, '题名', df['论文号'].map(lambda w: pp.at[w, '标题']))
    df.insert(3, '层级', df['论文号'].map(lambda w: sc.at[w, 'tier']))
    (d / '全文文本').mkdir(exist_ok=True)
    res, pages = [], []
    for _, r in df.iterrows():
        if not r['原句']:
            res.append(''); pages.append(''); continue
        s, spans = page_text(mf.at[r['论文号'], 'pdf'], d / '全文文本' / f"{r['论文号']}.txt")
        i = s.find(norm(r['原句']))
        res.append('✅' if i >= 0 else '❌')
        pages.append(next((p for a, b, p in spans if a <= i < b), '') if i >= 0 else '')
    df['原句核对'] = res
    df['页码'] = pages
    order = ['抑制', '权衡', '利用', '外加介体', '量化或校正', '只观察', '综述', '不计']
    df['_o'] = df['立场'].map(order.index)
    df = df.sort_values(['_o', '年份'], ascending=[True, False]).drop(columns='_o')
    cols = ['论文号', '年份', '题名', '层级', '立场', '做法', '条件变量', '无硫对照', '体系', '原句', '原句核对', '页码', '需人看', '备注', '人工确认']
    df['人工确认'] = ''
    df = df[cols]

    wb = Workbook(); ws = wb.active; ws.title = 'F逐篇复核表'
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9DEE6')
    for _, r in df.iterrows():
        ws.append([r[c] for c in cols])
        for c in ws[ws.max_row]:
            c.fill = PatternFill('solid', fgColor=FILL[r['立场']])
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, w in enumerate([13, 6, 40, 5, 9, 34, 20, 12, 14, 70, 7, 5, 6, 34, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    note = wb.create_sheet('说明')
    for line in [
        '范围：冻结 v2 论文表固态锂硫组中有本地全文的 828 篇。f_screen.py 按关键词分层：T1 50、T2 61 逐篇读；T3 163 只筛"本文实验"信号，补入 2 篇；T4 554 不读。',
        '立场：抑制=设法减少电解质在正极里的氧化还原；利用=把电解质（或其分解产物、含碘相）当活性物质或介体；权衡=同一工作里调一个条件、两种后果都看；外加介体=加入非电解质的介体；量化或校正=测或扣除电解质容量，不表态；只观察=提到电解质贡献容量，不作设计；综述=综述或展望；不计=负极侧、液态、非锂硫或只顺带提及。',
        '条件变量：该文改变或讨论、并认为影响结果的条件。',
        '无硫对照：是否用不含硫的电极（电解质加碳等）单独测电解质容量。',
        '原句核对：✅ 在本地全文逐字找到（忽略空白、连字符、连字和大小写）；❌ 没找到。页码是 PDF 页序。',
        '需人看：🔶 表示分类或说法有疑点，理由见备注。人工确认列留给你填。',
        '立场和做法是 Claude 的判断，不是作者原话；只有"原句"列是原文。',
    ]:
        note.append([line])
    note.column_dimensions['A'].width = 160
    wb.save(d / 'F_逐篇复核表.xlsx')

    inc = df[df['立场'] != '不计']
    orig = inc[~inc['立场'].isin(['综述'])]
    summ = {
        '读过': int(len(df)), '计入': int(len(inc)), '原创': int(len(orig)), '综述': int((inc['立场'] == '综述').sum()),
        '立场计数': inc['立场'].value_counts().to_dict(),
        '无硫对照_是': int(orig['无硫对照'].str.startswith('是').sum()),
        '原句': int((df['原句核对'] != '').sum()), '找到': int((df['原句核对'] == '✅').sum()),
        '需人看': int((df['需人看'] == '🔶').sum()),
        '年份立场': orig.groupby(['年份', '立场']).size().reset_index(name='n').to_dict('records'),
        '条件立场': [{'条件': c, '立场': r['立场']} for _, r in orig.iterrows() for c in r['条件变量'].split(';') if c and c != '—'],
    }
    (d / 'F_汇总.json').write_text(json.dumps(summ, ensure_ascii=False, indent=1), encoding='utf-8')
    for _, r in df[df['原句核对'] == '❌'].iterrows():
        print('未找到：', r['论文号'], r['原句'][:80])
    print({k: v for k, v in summ.items() if k not in ('年份立场', '条件立场')})


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
