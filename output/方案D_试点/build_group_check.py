"""方案D 组别核查（口径见 试点计划.md §D1）：60 篇随机样本逐篇判定 + K1–K5 定向排查 → 代码核对依据原句 → 复核表。

用法：python build_group_check.py <论文表.csv> <组别核查_抽样清单.csv> <组别核查材料目录> <输出.xlsx>
随机样本与定向排查分开统计，不合算总体错误率。不改 v2 文件。
"""
import math
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------- 60 篇逐篇判定 ----------------
# 序号: (文献类型, 液相状态, 液相依据, 负极, 实验配置, 研究对象, 含正极验证,
#        v2语义, 分类, 方案D归属, 主统计, 主统计原因, 第二层, 判据来源, 依据原句, 出处, 说明, 需人看原因)
# 出处：abs = 标题+v2 摘要；title = 只有标题；Gnn = 本地全文
R = {
 1: ('原创', '液态', '原文明示（碳酸酯电解液）', '锂金属', '全电池；对称电池', '正极宿主 + 正极界面（电解液添加剂）', '是',
     '不符合', '原规则误分：液态锂硫', '液态锂硫', '不纳入', '不属于本组', '相关', '摘要',
     'Stabilizing lithium-sulfur batteries in carbonate-based electrolytes', 'abs', 'v2 把"copolymer electrolyte additive"当作聚合物电解质', ''),
 2: ('原创（按标题）', '全固态', '标题（Li2S–P2S5 硫化物电解质）', '锂合金（Li–Si）', '全电池', '电解质氧化还原导致复合正极化学-机械失效', '是',
     '符合', '统计纳入：只有标题', '固态锂硫', '待核实', '只有标题，第一层统计前需补摘要（§4）', '相关', '标题',
     'comprising sulfur composite cathode and Li–Si alloy anode', 'title', '', '只有标题'),
 3: ('原创', '全固态', '原文明示', '锂金属', '全电池', '硫化物电解质液相合成，用作正极电解质、中间层、负极电解质', '是',
     '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
     'we realize a long-cycling all-solid-state lithium-sulfur battery', 'abs', '', ''),
 4: ('原创', '含液或凝胶', '原文明示（Li2S6 正极液）；离子液体依据配方推定', '锂金属', '全电池；对称电池', 'MOF 固态电解质抑制穿梭', '是',
     '不符合', '原规则误分：含液（固液混合）', '含液或凝胶锂硫', '不纳入', '含液（液态正极液）', '相关', '摘要+全文',
     'Lithium–sulfur cells with the Li‐IL/UIOSLi solid electrolyte and a Li 2 S 6 catholyte', 'abs',
     '全文：正极侧另加空白电解液，E/S = 5 μL/mg', '固液混合是否按 v2 本意算"不符合"'),
 5: ('原创', '全固态', '原文明示（正极内不加电解质）', '未写明', '全电池', '正极结构：粘结剂、3D 电极、高载量', '是',
     '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
     'without electrolyte in the cathode', 'abs', '', ''),
 6: ('原创', '含液或凝胶', '原文明示（离子液体为可流动的液态界面层）', '锂金属', '全电池；对称电池', '电极/石榴石电解质界面（离子液体修饰）', '是',
     '不符合', '原规则误分：含液（固液混合）', '含液或凝胶锂硫', '不纳入', '含液（离子液体界面层）', '相关', '全文',
     'plays a flowable and seamless interfacial layer in solid state Li-S batteries', 'G06', '', '固液混合是否按 v2 本意算"不符合"'),
 7: ('原创', '待核实', 'DMF 以配体固定在 Hofmann 框架中；是否仍有自由液态 DMF，原文未定量', '锂金属', '全电池；对称电池', '复合聚合物电解质（负极侧 DMF 分解）', '是',
     '符合', '体系边界：液相待核实；应用是通用锂金属电池', '待核实（通用锂金属电池电解质，全电池只配 SPAN）', '待核实', '液相待核实；应用归属待定', '不相关', '摘要+全文',
     'When paired with sulfurized polyacrylonitrile cathodes, the full cell exhibits', 'abs', '', 'Q7 液相与应用归属'),
 8: ('综述（课题组工作总结）', '液态为主', '原文明示（E/S、中间层）', '锂金属', '不适用', '锂硫各部件，聚合物电解质只占一节', '不适用',
     '不符合', '原规则误分：以液态为主的综述', '液态锂硫（综述，兼论聚合物电解质）', '不纳入', '综述；以液态为主', '相关（部分章节）', '摘要',
     'effectively reduce the electrolyte/sulfur ratio (E/S)', 'abs', '', '判断性：固态内容占比'),
 9: ('原创', '全固态', '原文明示', '锂合金（预锂化硅）', '全电池（含软包）', '电极应变协调、降低外压', '是',
     '符合', '体系边界：正极为 FeS2', '待核实（作者称固态锂硫，正极为金属硫化物 FeS2）', '待核实', 'FeS2 是否算硫基正极待定', '待核实', '摘要',
     'Using an FeS2 positive electrode and a prelithiated Si negative electrode as a representative system', 'abs', '', 'FeS2 是否算硫基正极'),
 10: ('原创（原位表征）', '全固态', '原文明示', '未写明', '材料测试（原位 TEM 纳米电池）', '正极反应机理', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'S/Li2S phase separation at nanoscale is observed for the first time in the solid-state Li–S battery', 'abs', '', ''),
 11: ('原创', '全固态', '原文明示', '锂合金（Li–In）', '全电池', '硫化物电解质合成与空气稳定性', '是',
      '符合', '', '固态锂硫', '纳入', '', '不相关', '全文',
      'hinders the exploitation of high energy all-solid-state lithium-sulfur batteries', 'G11', '', ''),
 12: ('原创', '含液或凝胶', '原文明示（电池中加 10 μL 液态电解液）', '锂金属', '全电池', '交联聚硅氧烷聚合物电解质', '是',
      '不符合', '原规则误分：含液（加液态电解液）', '含液或凝胶锂硫', '不纳入', '含液', '不相关', '全文',
      '10 μL of liquid electrolyte', 'G12', '', ''),
 13: ('原创', '全固态', '原文明示', '锂金属', '全电池', '复合正极：碳纤维骨架 + 原位包覆硫化物电解质', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '全文',
      'In the assembled all-solid-state Li–S battery with Li6PS5Cl electrolyte and lithium metal anode', 'G13', '', ''),
 14: ('原创', '全固态', '原文明示', '未写明', '全电池', '碳载体维度', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'we systematically investigated three representative carbon materials with distinct dimensionalities', 'abs', '', ''),
 15: ('原创', '含液或凝胶', '原文明示（电池中加 40 μL 液态电解液）', '锂金属', '全电池；对称电池', '交联聚醚聚合物电解质', '是',
      '不符合', '原规则误分：含液（加液态电解液）', '含液或凝胶锂硫', '不纳入', '含液', '不相关', '全文',
      '40 μL of 1 M LiTFSI (dioxolane: dimethoxyethane = 1:1 vol%) solution was added on top of P3', 'G15', '', ''),
 16: ('原创', '全固态', '原文明示', '未写明', '全电池', 'Li2S 正极氧化还原介质', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'A solid‐state cell with Li 2 S‐Li x In 2 S 3 cathode carries 82.35% capacity retention', 'abs', '', ''),
 17: ('综述', '全固态', '原文明示', '不适用', '不适用', '固态锂硫中的固态电解质', '不适用',
      '符合', '统计纳入：综述', '固态锂硫', '不纳入', '综述', '相关（部分章节）', '摘要',
      'In this review, a critical discussion on the key issues and problems of different types of SSEs', 'abs', '', ''),
 18: ('原创', '待核实', '原文：TGA 显示残留溶剂 < 8 wt%，作者称"严格说是固态或准固态"', '锂金属', '全电池；对称电池', '聚合物电解质与多硫化物的相互作用', '是',
      '符合', '体系边界：残留溶剂', '待核实（残留溶剂）', '待核实', '液相待核实', '相关', '摘要+全文',
      'still contain a small amount of solvent residue according to the TGA result (less than 8 wt %)', 'G18', '', '残留溶剂算不算含液'),
 19: ('原创', '全固态', '原文明示', '未写明', '全电池', '氧化还原活性正极电解质', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'If implemented as catholytes, this enables very fast sulfur conversion kinetics', 'abs', '', ''),
 20: ('会议论文集', '不适用', '', '不适用', '不适用', '一次电池、铅酸、镍铁等', '不适用',
      '不符合', '原规则误分：非锂硫或非电池', '其他体系（1980 年电源会议论文集）', '不纳入', '不属于本组', '不相关', '摘要',
      'lithium/sulfur dioxide primary batteries, lithium/oxychloride primary batteries', 'abs', '"oxychloride"触发卤化物和固态，"lithium/sulfur"触发锂硫', ''),
 21: ('原创', '全固态', '原文明示', '未写明', '全电池', '卤化物正极电解质', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'Li 3 YCl 5 I has been developed with a wide window of electrochemical stability for a catholyte', 'abs', '', ''),
 22: ('综述', '不适用', '', '不适用', '不适用', '原位 TEM，多种电池体系', '不适用',
      '不符合', '原规则误分：非锂硫或非电池', '其他（多体系方法综述）', '不纳入', '不属于本组', '不相关', '摘要',
      'including heavier alkali ion, lithium-sulfur, alkali-metal oxygen and all-solid-state batteries', 'abs', '"锂硫"和"全固态"是并列的两个体系', ''),
 23: ('原创', '全固态', '原文明示', '未写明', '全电池', '正极宿主（MoS2/C 混合导体，不加额外碳）', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'the S/MoS 2 /C electrode delivers excellent electrochemical performance', 'abs', '', ''),
 24: ('原创', '含液或凝胶', '依据配方推定（离子液体浸渍填料，填料:离子液体 = 3:4）', '锂金属', '全电池', 'PEO 复合电解质电导率', '是',
      '符合', '体系边界：含离子液体（Q7 推定含液）', '含液或凝胶锂硫（依据配方推定）', '不纳入', '含液（依据配方推定）', '不相关', '全文',
      'was mixed with the ionic liquid (IL) with a weight ratio of 3:4', 'G24', '作者称 composite solid electrolyte', 'Q7 推定是否成立'),
 25: ('原创', '全固态', '原文明示', '未写明', '全电池', '正极导电材料 + 单原子位点', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'Co–NC@S cathode delivers 1499 mAh g –1 at C/20', 'abs', '', ''),
 26: ('原创', '全固态', '原文明示（浇铸溶剂真空干燥去除）', '锂金属', '全电池', '聚合物电解质与 SPAN 正极的界面', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要+全文',
      'the all-solid-sate SPAN/PISE/Li batteries', 'abs', '', ''),
 27: ('原创（原位表征）', '全固态', '原文明示', '锂合金（InLi）', '全电池', '负极/电解质界面演化', '是',
      '符合', '', '固态锂硫', '纳入', '', '不相关', '标题+摘要',
      'in an all-solid-state Li–S cell by in operando synchrotron X-ray tomography', 'abs', '', ''),
 28: ('方法', '全固态', '原文明示', '锂合金（Li–In）', '全电池', '原位 SAXS/WAXS 电池设计；复合正极中 Li2S 分布', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'Applying this scanning SAXS/WAXS technique to all‐solid‐state lithium‐sulfur batteries', 'abs', '', ''),
 29: ('原创（按标题）', '全固态', '标题（电解质类别未知）', '未写明', '全电池（按标题）', '正极宿主', '是',
      '符合', '统计纳入：只有标题', '固态锂硫', '待核实', '只有标题，第一层统计前需补摘要（§4）', '相关', '标题',
      'enabling excellent cycling stability of all-solid-state lithium-sulfur batteries', 'title', '', '只有标题'),
 30: ('原创', '全固态', '原文明示（无溶剂捏合）', '锂金属', '全电池', '聚合物电解质加工与厚度', '是',
      '符合', '', '固态锂硫', '纳入', '', '不相关', '摘要',
      'a solvent-free polymer-based solid electrolyte was prepared by kneading', 'abs', '', ''),
 31: ('综述', '全固态', '', '不适用', '不适用', '锂硫用聚合物电解质', '不适用',
      '符合', '统计纳入：综述', '固态锂硫', '不纳入', '综述', '相关（部分章节）', '摘要',
      'This review critically examines recent advancements in polymer electrolytes for Li-S batteries', 'abs', '', ''),
 32: ('原创（原位表征）', '全固态', '原文明示（浇铸溶剂干燥去除）', '锂金属', '全电池', '多硫化物穿梭与界面演化（PEO/LLZTO）', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要+全文',
      'In situ monitoring of the interfacial processes in working all-solid-state lithium–sulfur batteries', 'abs', '', ''),
 33: ('综述', '不适用', '', '不适用', '不适用', '后锂电池技术', '不适用',
      '不符合', '原规则误分：非锂硫或非电池', '其他（多体系综述）', '不纳入', '不属于本组', '不相关', '摘要',
      'This paper presents a comprehensive exploration of emerging post-lithium battery technologies', 'abs', '', ''),
 34: ('原创', '含液或凝胶', '原文明示（正极侧界面加 8 μL 液态电解液）', '锂金属', '全电池；对称电池', '石榴石电解质制备（烧结助剂）', '是',
      '不符合', '原规则误分：含液（固液混合）', '含液或凝胶锂硫', '不纳入', '含液', '不相关', '全文',
      'which cathode electrode side interface was wetted by 8 μL liquid electrolyte', 'G34', '', '固液混合是否按 v2 本意算"不符合"'),
 35: ('综述', '不适用', '', '不适用', '不适用', '电动车电池综述', '不适用',
      '不符合', '原规则误分：非锂硫或非电池', '其他（多体系综述）', '不纳入', '不属于本组', '不相关', '摘要',
      'traces the historical evolution of battery technologies used in EVs', 'abs', '', ''),
 36: ('学位论文', '液态（推断）', '摘要：正极按常规粘结剂 + 碳制备，电解液未写', '未写明', '全电池（推断）', 'Li2S 纳米晶合成', '是',
      '不符合', '原规则误分：液态锂硫（推断）', '液态锂硫（推断）；固态电解质只是 Li2S 的用途之一', '不纳入', '不属于本组', '相关', '摘要',
      'Lithium sulfide (Li2S) nanocrystals (NCs) are critical materials used to make solid-state electrolytes and cathodes', 'abs', '', '判断性：液态为推断'),
 37: ('原创', '全固态', '原文明示', '锂合金（硅）', '全电池', 'Li2S 正极掺 Se', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'a full cell combining Li 2 Se 0.2 S 0.8 with a Si anode', 'abs', '', ''),
 38: ('原创（按标题）', '全固态', '标题（Li2S–P2S5 硫化物电解质）', '未写明', '全电池（按标题）', 'Li2S 正极碳包覆', '是',
      '符合', '统计纳入：只有标题', '固态锂硫', '待核实', '只有标题，第一层统计前需补摘要（§4）', '相关', '标题',
      'Carbon-coated Li2S cathode for improving the electrochemical properties of an all-solid-state lithium-sulfur battery', 'title', '', '只有标题'),
 39: ('原创', '液态', '原文明示（DOL/DME）', '锂金属', '全电池；对称电池', '锂负极人工混合固态电解质层', '是',
      '不符合', '原规则误分：液态锂硫', '其他：液态锂金属负极研究（配 LFP 和硫正极）', '不纳入', '不属于本组', '不相关', '全文',
      'polysulfides are soluble in the liquid elec- trolyte of DOL and DME', 'G39', '固态电解质只作负极人工层', ''),
 40: ('原创', '全固态', '原文称"全固态"，电解质未写', '未写明', '全电池', '3D 碳布正极', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要（韩文）',
      '전고체 리튬-황 전지는', 'abs', '', ''),
 41: ('原创', '全固态', '原文明示（PEO 基全固态）', '锂金属', '全电池', '正极中 g-C3N4（吸附与催化）', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '全文',
      'Poly(ethylene oxide)-based polymer all-solid-state lithium-sulfur batteries (ASSLSBs)', 'G41', '', ''),
 42: ('原创（材料基础）', '不适用', '', '不适用', '材料测试', 'Li2S 缺陷化学与离子传输', '否',
      '不符合', '原规则误分：非锂硫或非电池', '其他（材料基础研究，未明示固态锂硫）', '不纳入', '不属于本组', '相关（可作支撑材料）', '摘要',
      'An application of these results to passivation layers in solid state batteries is also discussed', 'abs', '', '判断性：是否算锂硫研究'),
 43: ('原创', '全固态', '原文明示（PEO 固态聚合物电解质）', '锂金属', '全电池', '正极 GO–PEG 改性', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要+全文',
      'basing Li–S batteries on a solid electrolyte (SE) instead of a liquid electrolyte', 'G43', '', ''),
 44: ('原创', '全固态', '原文明示', '未写明', '全电池', '正极氧化还原介质 WS2', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'The hybrid cathode in ASSLSB exhibits outstanding capacity retention', 'abs', '', ''),
 45: ('综述', '全固态', '', '不适用', '不适用', '复合正极中的载流子传输', '不适用',
      '符合', '统计纳入：综述', '固态锂硫', '不纳入', '综述', '相关', '摘要',
      'we provide research progresses on facilitating transport of charge carriers in composite cathode of SSLSBs', 'abs', '', ''),
 46: ('原创', '待核实', '原文称"hybrid Li−S battery"，正极侧组成正文未写', '锂金属', '全电池；对称电池', 'LLZO 烧结工艺', '是',
      '证据不足', '证据不足', '证据不足（LLZO 制备，锂硫只是演示）', '待核实', '液相待核实', '不相关', '全文',
      'The hybrid Li− S battery delivered an initial discharge specific capacity', 'G46', 'SI 未查', '证据不足'),
 47: ('原创', '全固态', '原文明示（浇铸溶剂挥发后加热去除）', '锂金属', '全电池', 'PEO 电解质填料（电导率）', '是',
      '符合', '', '固态锂硫', '纳入', '', '不相关', '摘要+全文',
      'blended with polyethylene oxide to prepare a solid polymer electrolyte', 'abs', '', ''),
 48: ('原创', '液态', '原文明示（DME:DOL）', '未写明', '全电池', '粘结剂堵塞碳孔隙', '是',
      '不符合', '原规则误分：液态锂硫', '液态锂硫', '不纳入', '不属于本组', '相关', '摘要',
      'Because of the low swellability of PVdF in dimethoxyethane:dioxolane (DME:DOL)-based electrolytes', 'abs', 'v2 由"PVdF-HFP"触发聚合物/固态', ''),
 49: ('学位论文', '液态为主；复合聚合物电解质部分状态不明', '', '未写明', '全电池', '硫正极结构 + 复合聚合物电解质', '是',
      '证据不足', '证据不足', '液态锂硫为主（学位论文）；复合聚合物电解质部分待核实', '不纳入', '学位论文；固态部分证据不足', '相关', '摘要',
      'This PhD dissertation overcomes these limitations by engineering electron conducting novel cathode architectures', 'abs', '', '证据不足'),
 50: ('综述', '液态为主', '', '不适用', '不适用', '锂硫新构型（梯度正极、隔膜、固态电解质、负极保护）', '不适用',
      '不符合', '原规则误分：以液态为主的综述', '液态锂硫（综述，固态电解质为一节）', '不纳入', '综述', '相关（部分章节）', '摘要',
      'including hierarchical gradient cathodes, modified separators, solid-state electrolytes and lithium anode protection', 'abs', '', '判断性：固态内容占比'),
 51: ('封面介绍（非研究论文）', '含液或凝胶', '原文明示（凝胶电解质）', '不适用', '不适用', '印刷双极锂硫电池', '不适用',
      '不符合', '原规则误分：含液（凝胶）', '含液或凝胶锂硫（封面介绍）', '不纳入', '非研究论文；凝胶', '不相关', '摘要',
      'two thermodynamically immiscible and nonflammable gel electrolytes are used', 'abs', '作者标题写 All-Solid-State', ''),
 52: ('原创', '液态', '原文明示（凝胶聚合物只作中间层）', '未写明', '全电池', '3D 正极 + 凝胶中间层', '是',
      '不符合', '原规则误分：液态锂硫', '液态锂硫', '不纳入', '不属于本组', '相关', '摘要',
      'CNF additive is used as an interlayer to trap polysulfides', 'abs', '与批1 #5 重合，判定一致', ''),
 53: ('原创', '全固态', '原文明示', '锂金属', '全电池；对称电池', '硫化物电解质改性抑制枝晶', '是',
      '符合', '体系边界：同时测 NCM811 与硫正极', '固态锂硫 + 非硫固态（两种正极都测）', '纳入', '双归属，计数方式待定', '不相关', '摘要',
      'support stable cycling of both NCM811 and sulfur cathodes paired with Li metal', 'abs', '', '双归属怎么计'),
 54: ('原创', '含液或凝胶', '原文明示（液态 Li2S8 正极液、液态锂汞齐负极）', '液态合金（锂汞齐）', '全电池；对称电池', '液态电极/固态电解质界面', '是',
      '不符合', '原规则误分：含液（固液混合）', '含液或凝胶锂硫', '不纳入', '含液', '相关', '全文',
      'Liquid Li2S8 polysulfide catholyte can afford a high active material loading', 'G54', '', '固液混合是否按 v2 本意算"不符合"'),
 55: ('原创', '全固态', '原文明示（DMF 浇铸，60 °C 干燥 24 h）', '锂金属', '全电池；对称电池', '复合电解质填料（离子传导与硫转化）', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '全文',
      'cast onto a polytetrafluoroethylene (PTFE) mold and dried at 60', 'G55', '浇铸溶剂有去除步骤，按 Q7 不据此判含液', ''),
 56: ('原创', '全固态', '原文明示', '锂金属', '全电池', '纳米限域 LiBH4 电解质', '是',
      '符合', '', '固态锂硫', '纳入', '', '不相关', '摘要',
      'based on nano-confined LiBH 4 in mesoporous silica as solid electrolytes', 'abs', '真正的硼氢化物电解质', ''),
 57: ('原创', '全固态', '原文明示', 'LTO（对电极）', '全电池', 'Li2S 基正极体积变化与堆压', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '摘要',
      'by using Li 4 Ti 5 O 12, which is well-known to undergo almost no volume change during cycling, as the counter electrode', 'abs', '', ''),
 58: ('综述', '全固态', '', '不适用', '不适用', '固态锂硫的高电导固态电解质', '不适用',
      '符合', '统计纳入：综述', '固态锂硫', '不纳入', '综述', '相关（部分章节）', '标题',
      'Recent Progress in All-Solid-State Lithium−Sulfur Batteries Using High Li-Ion Conductive Solid Electrolytes', 'title', '', ''),
 59: ('综述', '全固态', '', '不适用', '不适用', '固态锂硫实用化', '不适用',
      '符合', '统计纳入：综述', '固态锂硫', '不纳入', '综述', '相关（部分章节）', '摘要',
      'Here, this Review examines their recent advances', 'abs', '', ''),
 60: ('原创', '全固态', '原文明示', '锂金属', '全电池', '石墨烯/硫复合正极合成路线', '是',
      '符合', '', '固态锂硫', '纳入', '', '相关', '全文',
      'were utilized as anode and solid electrolyte, respectively', 'G60', '', ''),
}
# ---------------- 第二轮裁定（2026-10-18 晚，用户转来的评审意见；对照课题方案附录 D 冻结规则）----------------
# 原则：只有违背 v2 分组本意且证据充分的，才记"不符合（已确认）"；附录 D 未覆盖的固液混合等记"边界未明确"；
# 多体系或以液态为主的综述、推断记"判断性"。区间只反映抽样不确定性，不含判定不确定性。
# 序号: {字段: 新值}
FIX = {
    1: {'v2 语义归组': '不符合（已确认）', '分类': '已确认误分'},
    4: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态电解质 + 液态正极液）', '需人看的原因': ''},
    6: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态电解质 + 离子液体界面层）', '需人看的原因': ''},
    8: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：以液态为主的综述'},
    9: {'分类': '体系边界：正极为 FeS2（D1 按标题判，标题写锂硫，不排除）', '方案D 应用归属': '金属硫化物转化正极（单列；有相关功能证据时作支撑材料）',
        '主统计纳入': '不纳入', '主统计原因': '单列，不进硫正极主组', '需人看的原因': ''},
    12: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态聚合物电解质 + 电池中加液态电解液）'},
    15: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态聚合物电解质 + 电池中加液态电解液）'},
    18: {'方案D 应用归属': '待核实（原文说"若干制备样品"残留溶剂 < 8 wt%，需对应到电池所用样品）', '主统计原因': '暂不进全固态主统计；< 8 wt% 不作合格阈值',
         '需人看的原因': '残留溶剂对应到哪个样品'},
    20: {'v2 语义归组': '不符合（已确认）', '分类': '已确认误分'},
    22: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：多体系综述', '需人看的原因': '判断性：多体系综述'},
    24: {'液相依据': '依据配方推定：离子液体与填料按 4:3（质量）制成 TMCs-IL，后经 120 ℃ 真空处理 12 h；不是自由液相比例的直接测量，也不能据真空处理认定离子液体已除去',
         '需人看的原因': ''},
    33: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：多体系综述', '需人看的原因': '判断性：多体系综述'},
    34: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态电解质 + 正极侧液态电解液）', '需人看的原因': ''},
    35: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：多体系综述', '需人看的原因': '判断性：多体系综述'},
    36: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：液态为推断'},
    39: {'v2 语义归组': '不符合（已确认）', '分类': '已确认误分'},
    42: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：是否算锂硫研究'},
    48: {'v2 语义归组': '不符合（已确认）', '分类': '已确认误分'},
    50: {'v2 语义归组': '不符合（判断性）', '分类': '判断性：以液态为主的综述'},
    51: {'v2 语义归组': '边界未明确', '分类': '冻结规则按标题自称判固态（D3.1），方案D 按凝胶排除'},
    52: {'v2 语义归组': '不符合（已确认）', '分类': '已确认误分'},
    53: {'方案D 应用归属': '固态锂硫 + 非硫固态（双归属：论文主表一篇，体系配置两条）',
         '主统计原因': '双归属：各组每篇最多计一次，只把有证据属于该体系的标签计入该组；组间有共享论文', '需人看的原因': ''},
    54: {'v2 语义归组': '边界未明确', '分类': '旧规则边界未明确，本次明确排除（固态电解质 + 液态正极液、液态合金负极）', '需人看的原因': ''},
}
RULE = {  # 规则层面原因（修规则时用）
    1: '"copolymer electrolyte additive"子串触发聚合物电解质（规则未按文字执行）',
    4: 'D3.2 字面命中（写了 solid electrolyte）；附录 D 未覆盖液态正极液', 6: 'D3.2 字面命中（石榴石）；附录 D 未覆盖液态界面层',
    8: '摘要并列提到锂硫与固态电解质（D8 已知局限）', 50: '同 #8',
    9: 'D1 只按标题排除金属硫化物；标题写锂硫，规则下不排除',
    12: 'D3.2 字面命中（solid polymer electrolyte）；加液只写在方法里', 15: '同 #12',
    20: '"lithium/sulfur dioxide"触发锂硫、"oxychloride"触发固态（规则未按文字执行）',
    22: '多体系综述并列提到锂硫与全固态（D8 已知局限）', 33: '同 #22', 35: '同 #22',
    24: 'D5 复合电解质（聚合物 + 无机物）；离子液体不在规则视野内',
    34: 'D3.2 字面命中（石榴石）；附录 D 未覆盖正极侧液态电解液',
    36: '"Li2S cathodes"与"solid-state electrolytes"并列（D8 已知局限）',
    39: 'D3.2 字面命中（固态电解质作负极人工层），规则看不到用途（D8 已知局限）',
    42: '"Li–S cathodes"与"solid state batteries"并列（D8 已知局限）',
    48: '"PVdF-HFP"粘结剂触发聚合物电解质（规则未按文字执行）',
    51: 'D3.1 明确规定标题自称固态时按固态判',
    52: 'PVDF-HFP 凝胶中间层：摘要提到 gel，D3.3 要求"没提凝胶"，却仍判固态（规则未按文字执行）',
    54: 'D3.2 字面命中（solid-state electrolytes）；附录 D 未覆盖液态电极',
}

COLS = ['文献类型', '液相状态', '液相依据', '负极类型', '实验配置', '研究对象', '含正极的电池验证',
        'v2 语义归组', '分类', '方案D 应用归属', '主统计纳入', '主统计原因', '第二层相关性', '判据来源',
        '依据原句', '出处', '说明', '需人看的原因']

# ---------------- 定向排查：逐篇判定代码（按候选文件行序） ----------------
K2A = ('R L R L Q S L R O N|N O R R R L L R ? L|S L L Q L O O O R R|L L L R L L L L ? L|R Q L S L O O O N O|'
       'O R L L L R Q S S N|O Q R R ? R Q L L ?|L L R L O O O L L R|Q R L L L L L L O N|? O O ? L L R L S L|'
       'L Q O Q N O L Q L Q|? ? L Q Q L L O L N|L ? R R ? L L L Q L|L O Q L L L O Q R Q|Q L L L L L L L L R|S N S L L').replace('|', ' ').split()
K2B = ('Q S Q ? S S S R S ? O N R Q R S Q S S O O ? ? S S S S Q S Q ? ? Q ? S S ? S S L L N Q S S S S ? O S S S R S O Q S R N R Q Q S ? S S O ? S ? O R '
       'O S S').split()
K3S = 'R S 双 N N O ? O N N N S N O N O O L S N N O O O'.split()
K3L = ('OK OK OK OK 非锂 非锂 OK OK OK OK ? OK OK 非锂 负极 负极 负极 OK OK OK 其他 OK 非电池 非锂 非锂 OK OK OK 负极 OK 负极 负极 ? OK OK OK ? OK OK 其他 '
       '? 负极 负极 负极 ? 非锂 误触发 OK OK 负极 ? 非锂 非锂 综述 非电池 OK 综述 ? OK OK OK 非电池 非锂 OK OK OK 误触发 OK 误触发 误触发 误触发 OK 综述 非锂 OK 综述 其他').split()
K4 = ('水 非锂 水 非水 非水 非水 非水 非锂 ? 水 非电池 非锂 非锂 水 非水 非水 水 非锂 水 非水 非水 水 非锂 非水 非水 非水 非水 非锂 非锂 非锂 水 非水 非锂 液 非锂 非锂 非锂 ? 水 非锂 '
      '水 非锂 非锂 水 非水 水 非锂 非锂 非水 水 ? 非锂 非锂 非水 非锂 非锂 液 水 非水 水 非锂 非水 非水 非水 水 水 水 非水 非水 非锂 水 水 水 水 非水 ? 非锂 水 非锂 水 '
      '水 水 水 非水 水 非水 ? 水 液 非水 液 水 水 非锂 水 水 非水 非水 水 水 水 非水 水 水 非锂 液 水 水 水 水 水 非水 水 水 水 水 水 水 ? 非水 水 水 水 非水 水 水 非水 非水').split()
K5 = ('合规 合规 合规 合规 合规 合规 非锂 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 非锂 一次 ? 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 合规 一次 非锂 合规 合规 合规 合规 '
      '非锂 合规 合规 合规 一次 合规 合规 合规 合规 非锂 非锂 合规 合规 合规 合规 ? 一次 非锂 非锂 合规 一次 一次 误触发 S 非锂 误触发 一次 OK OK 误触发 OK ? 误触发 OK ? 误触发 误触发 误触发').split()
K1 = {  # 编号: (判定, 说明)
    'W4414998051': ('误触发', '还原剂 NaBH4 制备正极碳；液态锂硫'),
    'W2779656978': ('误触发', '还原剂 NaBH4 还原氧化石墨烯；液态锂硫（批1 #4）'),
    'W2109705256': ('误触发', '硼氢化锂只是合成锂盐的前驱体；液态电解液'),
    'W7161822004': ('误触发', '还原剂 NaBH4；室温钠硫电池，非锂'),
    'W1981619380': ('误触发', '还原剂 LiEt3BH 制备 Li2S；液态锂硫正极，被分进固态锂电池'),
    'W4414748693': ('组别问题', '标签对（复合氢化物），但作液态电解液添加剂；液态锂硫'),
    'W4387003065': ('非锂', '钠离子固态电解质'), 'W4401126679': ('非锂', '全固态钠电池综述'),
    'W4407673921': ('非锂', '镁硼氢化物电解质'), 'W4293574805': ('非锂', '镁硼氢化物电解质'),
    'W3080472362': ('非锂', '全固态镁电池'), 'W3012936885': ('非锂', '镁离子传导'),
    'W4289334449': ('非锂', '钙电池，THF 液态电解液'),
}
CODE = {'L': '液态锂硫误收', 'Q': '含液或固液混合（应归准固态或单列）', 'R': '以液态为主或多体系的锂硫综述', 'O': '非锂硫或非电池',
        'N': '非硫固态（无硫正极）', 'S': '确属固态锂硫', '?': '证据不足', '双': '固态锂硫 + 非硫固态双归属',
        'OK': '确属液态锂硫', '负极': '锂金属负极研究，含 Li‖S 全电池（不算误收）', '非锂': '非锂体系', '非电池': '不是电池研究',
        '误触发': '关键词误触发', '综述': '多体系综述', '其他': '其他电池体系（热电池、水系等）',
        '水': '水系或含水锂体系（含水系锂空气）', '非水': '与水系电池无关（"aqueous"误命中、综述或本组合规），本项不判',
        '液': '非水液态电解液', '一次': '一次电池（Li/SOCl2、Li/CFx、热电池等）', '合规': '卤氧化物等固态电解质，标签合规'}
MISFILED = {'L', 'Q', 'R', 'O', 'N', '非锂', '非电池', '误触发', '综述', '其他', '组别问题'}

# ---------------- 敏感性情景：固态锂硫组标记清单的证据等级（第二轮裁定）----------------
# 明示 = 摘要明写液态电池构型（隔膜改性、液态电解液、电解液添加剂、E/S 等）或明写不是锂硫电池；
# 触发词误用 = 唯一的"固态"证据是粘结剂、还原剂、催化剂等误触发，冻结规则按文字执行本就判不成固态；
# 推断 / 判断性 = 不进情景 A、B，单独列出。新组：液态锂硫 / 固态锂电池 / 准固态锂硫 / 不进组（引用它的记池外）
LV = {  # (清单, 行号): (证据等级, 新组)
    **{('K2A', i): ('明示', '液态锂硫') for i in [2, 4, 7, 16, 23, 25, 32, 33, 35, 36, 37, 38, 40, 45, 53, 54, 55, 72, 74, 85, 86, 88, 95, 96, 98, 101,
                                               107, 109, 116, 117, 126, 127, 130, 131, 136, 143, 144, 145, 146, 147]},
    **{('K2A', i): ('触发词误用', '液态锂硫') for i in [17, 68, 79, 83, 84, 87, 113, 128, 134, 142, 148, 149, 154, 155]},
    **{('K2A', i): ('推断', '液态锂硫') for i in [20, 22, 31, 43, 69, 71, 78, 100, 119, 121, 135]},
    ('K2A', 12): ('明示', '不进组'), ('K2A', 89): ('明示', '不进组'), ('K2A', 137): ('明示', '不进组'), ('K2A', 51): ('推断', '不进组'),
    ('K2A', 11): ('触发词误用', '固态锂电池'), ('K2A', 105): ('明示', '固态锂电池'), ('K2A', 120): ('明示', '固态锂电池'),
    ('K2A', 10): ('推断', '固态锂电池'), ('K2A', 49): ('推断', '固态锂电池'),
    ('K2B', 40): ('推断', '液态锂硫'), ('K2B', 41): ('明示', '液态锂硫'), ('K2B', 59): ('明示', '固态锂电池'), ('K2B', 12): ('推断', '固态锂电池'),
    ('K3S', 18): ('明示', '液态锂硫'), ('K3S', 6): ('明示', '不进组'), ('K3S', 8): ('明示', '不进组'), ('K3S', 20): ('明示', '固态锂电池'),
    **{('K3S', i): ('推断', '固态锂电池') for i in [4, 10, 11, 13, 21]},
    ('K5', 63): ('触发词误用', '液态锂硫'), ('K5', 65): ('明示', '不进组'), ('K5', 66): ('明示', '不进组'),
}
K1_LV = {'W4414998051': ('触发词误用', '液态锂硫'), 'W2779656978': ('明示', '液态锂硫'), 'W2109705256': ('明示', '液态锂硫'), 'W4414748693': ('明示', '液态锂硫')}
SAMPLE_SC = {  # 随机样本：序号 → (情景, 新组)
    1: ('A', '液态锂硫'), 20: ('A', '不进组'), 39: ('A', '液态锂硫'), 48: ('A', '液态锂硫'), 52: ('A', '液态锂硫'),
    **{i: ('B', '准固态锂硫') for i in [4, 6, 12, 15, 34, 51, 54, 24]}, 9: ('B', '不进组'),
}


def norm(s):
    s = unicodedata.normalize('NFKC', str(s))
    return re.sub(r'[\s\-‐‑‒–—―−]', '', s).lower()


def wilson(k, n, z=1.96):
    if n == 0:
        return (float('nan'),) * 2
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return c - h, c + h


def pct(k, n):
    lo, hi = wilson(k, n)
    return f'{k}/{n}（{k / n:.0%}；95% 区间 {lo:.0%}–{hi:.0%}）'


def main(paper_csv, sample_csv, mat_dir, out):
    mat = Path(mat_dir)
    d = pd.read_csv(paper_csv, dtype=str, low_memory=False).fillna('')
    idx = d.set_index('编号')
    s = pd.read_csv(sample_csv, dtype=str).fillna('')

    rows = []
    for _, r in s.iterrows():
        n = int(r['序号'])
        v = dict(zip(COLS, R[n]))
        v.update(FIX.get(n, {}))
        src = v['出处']
        if src == 'abs':
            text = idx.at[r['论文号'], '标题'] + ' ' + idx.at[r['论文号'], '摘要']
        elif src == 'title':
            text = idx.at[r['论文号'], '标题']
        else:
            text = (mat / '全文' / f'{src}.txt').read_text(encoding='utf-8')
        found = norm(v['依据原句']) in norm(text)
        mark = '❌' if not found else ('🔶' if v['需人看的原因'] else '✅')
        rows.append({'序号': n, '论文号': r['论文号'], 'DOI': r['DOI'], '年份': r['年份'], '标题': r['标题'], '来源': r['来源'],
                     'v2 电解质类型': r['v2 电解质类型'], **{k: v[k] for k in COLS[:-3]},
                     '出处': {'abs': '标题+摘要', 'title': '标题'}.get(src, f'本地全文 {src}'),
                     '原句核对': '找到' if found else '未找到', '标记': mark, '需人看的原因': v['需人看的原因'],
                     '规则层面原因': RULE.get(n, ''), '说明': v['说明'], '复核 ✍': ''})
    g = pd.DataFrame(rows)

    # 定向排查明细
    kr = []
    def add(item, grp, cand, codes, note=''):
        assert len(cand) == len(codes), (item, len(cand), len(codes))
        for (_, c), code in zip(cand.iterrows(), codes):
            kr.append({'排查项': item, '组': c.get('分组', grp) or grp, '论文号': c['编号'], 'DOI': c.get('DOI', ''), '年份': c.get('年份', ''),
                       '标题': c['标题'], '判定': CODE.get(code, code), '误收': '是' if code in MISFILED else ('待核实' if code == '?' else '否'), '复核 ✍': ''})
    b = d[d['电解质类型'] == '硼氢化物']
    for k, (code, note) in K1.items():
        c = b[b['编号'] == k].iloc[0]
        kr.append({'排查项': 'K1 硼氢化物标签误触发', '组': c['分组'], '论文号': k, 'DOI': c['DOI'], '年份': c['年份'], '标题': c['标题'],
                   '判定': f'{CODE.get(code, code)}：{note}', '误收': '是', '复核 ✍': ''})
    add('K2 固态锂硫组·液态关键词（无固态字样）', '固态锂硫', pd.read_csv(mat / 'K2_候选.csv', dtype=str).fillna(''), K2A)
    add('K2 固态锂硫组·液态关键词（有固态字样）', '固态锂硫', pd.read_csv(mat / 'K2_候选_有固态词.csv', dtype=str).fillna(''), K2B)
    add('K3 固态锂硫组·非锂硫或非电池', '固态锂硫', pd.read_csv(mat / 'K3_固态锂硫.csv', dtype=str).fillna(''), K3S)
    add('K3 液态锂硫组·非锂硫或非电池', '液态锂硫', pd.read_csv(mat / 'K3_液态锂硫.csv', dtype=str).fillna(''), K3L)
    add('K4 固态锂电池组·水系电池（按标题粗分）', '固态锂电池', pd.read_csv(mat / 'K4_精筛.csv', dtype=str).fillna(''), K4)
    add('K5 一次电池关键词', '', pd.read_csv(mat / 'K5_一次电池.csv', dtype=str).fillna(''), K5)
    k = pd.DataFrame(kr)

    # 敏感性情景清单（只针对 v2 固态锂硫组；随机样本的逐篇裁定优先于定向排查）
    lists = {'K2A': ('K2 固态锂硫组·液态关键词（无固态字样）', K2A), 'K2B': ('K2 固态锂硫组·液态关键词（有固态字样）', K2B),
             'K3S': ('K3 固态锂硫组·非锂硫或非电池', K3S), 'K5': ('K5 一次电池关键词', K5)}
    pp = d.set_index('编号')
    sc = {}
    for key, (item, codes) in lists.items():
        sub = k[k['排查项'] == item].reset_index(drop=True)
        for i, row in sub.iterrows():
            if row['组'] != '固态锂硫':
                continue
            code = codes[i]
            lvl, new = LV.get((key, i + 1), (None, None))
            if lvl is None:
                lvl = {'Q': '边界', 'S': '无改动', '双': '无改动', '?': '证据不足'}.get(code, '判断性')
                new = {'Q': '准固态锂硫'}.get(code, '')
            sc.setdefault(row['论文号'], []).append((f'{key}#{i + 1}', row['判定'], lvl, new))
    for kid, (lvl, new) in K1_LV.items():
        sc.setdefault(kid, []).append(('K1', '关键词误触发', lvl, new))
    rank = {'明示': 0, '触发词误用': 1, '边界': 2, '推断': 3, '判断性': 4, '证据不足': 5, '无改动': 6}
    rows_sc = []
    smp = {r['论文号']: int(r['序号']) for _, r in g.iterrows()}
    for pid in set(sc) | set(smp):
        hits = sorted(sc.get(pid, []), key=lambda x: rank[x[2]])
        if pid in smp:
            n0 = smp[pid]
            sgr = g[g['序号'] == n0].iloc[0]
            how, new = SAMPLE_SC.get(n0, (None, ''))
            lvl = {'A': '随机样本：已确认', 'B': '随机样本：边界'}.get(how) or (
                '随机样本：体系边界，不改' if sgr['分类'].startswith('体系边界') else '随机样本：' + sgr['v2 语义归组'])
            a = how == 'A'
            b = how in ('A', 'B')
            basis = sgr['分类']
            found = '随机样本 #%d' % n0 + ('；' + '、'.join(h[0] for h in hits) if hits else '')
        else:
            top = hits[0]
            lvl, new = top[2], top[3]
            a = lvl in ('明示', '触发词误用')
            b = a or lvl == '边界'
            basis = top[1]
            found = '、'.join(h[0] for h in hits)
        if pid in smp and not (a or b):
            sgr = g[g['序号'] == smp[pid]].iloc[0]
            if sgr['v2 语义归组'] == '符合' and not sgr['分类'].startswith('体系边界'):
                continue  # 随机样本中确属本组、也不涉及边界的不列
        p = pp.loc[pid]
        rows_sc.append({'论文号': pid, '标题': p['标题'], '年份': p['年份'], '来源': p['来源'], 'v2 分组': p['分组'], '发现途径': found,
                        '判定': basis, '证据等级': lvl, '情景 A（已确认错误修正）': new if a else '', '情景 B（严格全固态边界）': new if b else '',
                        '改动对象': ('引用方样本 + 被引方标签' if p['来源'] == '检索' else '被引方标签（补池论文不作引用方）') if (a or b) else '',
                        '复核 ✍': ''})
    scdf = pd.DataFrame(rows_sc).sort_values(['情景 A（已确认错误修正）', '情景 B（严格全固态边界）', '证据等级'], ascending=[False, False, True])
    scdf[['论文号', '情景 A（已确认错误修正）', '情景 B（严格全固态边界）']].to_csv(Path(out).with_name('敏感性情景清单.csv'), index=False, encoding='utf-8-sig')

    # 汇总
    n = len(g)
    S = []
    S.append(['一、固态锂硫组随机 60 篇（简单随机，种子 20261018）', ''])
    S.append(['抽样框与分析一引用方', '抽样框 = v2 固态锂硫组 1,181 篇（检索 1,146 + 补池 35）；分析一引用方只用检索来的 1,146 篇。'
              '60 篇恰好全是检索来源，可视为引用方的简单随机样本；但比例不能直接当成分析一的错误率，还要看各篇的引用数'])
    S.append(['计数说明', '只报分类计数，不作错误率估计，也不表示审计通过（试点计划 §E）。判断性、边界未明确、证据不足各类单独列出'])
    for lab in ['符合', '不符合（已确认）', '不符合（判断性）', '边界未明确', '证据不足']:
        sub = g[g['v2 语义归组'] == lab]
        S.append([f'v2 语义归组：{lab}', f'{len(sub)}/{n}' + ('；' + '、'.join('#' + str(x) for x in sub['序号']) if lab != '符合' else '')])
    bnd = g[(g['v2 语义归组'] == '符合') & g['分类'].str.startswith('体系边界')]
    S.append(['v2 符合、但方案D 下单列 / 待核实 / 不纳入（体系边界调整）', f'{len(bnd)} 篇：' + '；'.join(f"#{r['序号']} {r['方案D 应用归属']}" for _, r in bnd.iterrows())])
    for lab in ['纳入', '待核实', '不纳入']:
        S.append([f'方案D 主统计：{lab}', f"{(g['主统计纳入'] == lab).sum()}/{n}"])
    why = g[g['主统计纳入'] == '不纳入']['主统计原因'].str.split('；').str[0].str.replace(r'（.*', '', regex=True)
    S.append(['　不纳入原因', '；'.join(f'{a} {b}' for a, b in why.value_counts().items())])
    inc = g[g['主统计纳入'] == '纳入']
    S.append(['主统计纳入的论文中第二层相关', f"{(inc['第二层相关性'] == '相关').sum()}/{len(inc)}"])
    S.append(['全部 60 篇第二层相关性', '；'.join(f'{a} {b}' for a, b in g['第二层相关性'].str.replace(r'（.*', '', regex=True).value_counts().items())])
    S.append(['依据原句核对（只说明原句能在来源中找到，不说明判定正确）', f"找到 {(g['原句核对'] == '找到').sum()}/{n}"])
    S.append(['标记', '；'.join(f'{a} {b}' for a, b in g['标记'].value_counts().items())])
    S.append(['', ''])
    S.append(['二、定向排查标记清单（关键词圈出候选后逐篇看标题和摘要；每项单独报，不与随机样本合算，也不是错误率下限）', ''])
    for item, grp in k.groupby('排查项', sort=False):
        vc = grp['判定'].str.replace(r'：.*', '', regex=True).value_counts()
        S.append([f'{item}：候选 {len(grp)} 篇', '；'.join(f'{a} {b}' for a, b in vc.items())])
    sl = k[(k['组'] == '固态锂硫') & (k['误收'] == '是')]['论文号'].unique()
    S.append(['固态锂硫组：被标记的论文（去重）', f'{len(sl)} 篇；含多体系综述、固液混合等判断性或边界类别，不整体当作已确认误收'])
    ll = k[(k['组'] == '液态锂硫') & (k['误收'] == '是')]['论文号'].unique()
    S.append(['液态锂硫组：被标记的论文（去重）', f'{len(ll)} 篇；只查了非锂硫/非电池和一次电池两类'])
    S.append(['K4（固态锂电池组水系，按标题粗分）', '只作线索，不作删除或改标签依据'])
    S.append(['', ''])
    S.append(['三、敏感性情景清单（固态锂硫组；见"敏感性情景清单"页）', ''])
    S.append(['情景 A：已确认错误修正', f"{(scdf['情景 A（已确认错误修正）'] != '').sum()} 篇：" + '；'.join(
        f'{a} {b}' for a, b in scdf[scdf['情景 A（已确认错误修正）'] != '']['情景 A（已确认错误修正）'].value_counts().items()) +
        '。只收随机样本已确认 + 定向排查中证据等级为"明示"或"触发词误用"的'])
    S.append(['情景 B：严格全固态边界', f"{(scdf['情景 B（严格全固态边界）'] != '').sum()} 篇：" + '；'.join(
        f'{a} {b}' for a, b in scdf[scdf['情景 B（严格全固态边界）'] != '']['情景 B（严格全固态边界）'].value_counts().items()) +
        '。A 之外再改固液混合、含液推定等边界记录（#9 金属硫化物正极移出）'])
    S.append(['两个情景都不改的标记', '；'.join(f'{a} {b}' for a, b in scdf[scdf['情景 B（严格全固态边界）'] == '']['证据等级'].value_counts().items())])
    S.append(['改动方式', '只改分组标签，不删引用关系：被改的检索论文移出引用方，同时在被引方改记新组；改成"不进组"的，引用它的记池外。'
              '其余参数与冻结版一致，结果另存，不覆盖冻结文件'])
    S.append(['', ''])
    S.append(['四、误触发类型（修规则时用）', ''])
    for a, b in [
        ('固态材料在液态电池里的用法触发"固态"', 'PEO、PVdF-HFP 作粘结剂或中间层；LLTO、LATP、LLZO、Li3PS4 作隔膜涂层、中间层或人工 SEI；"copolymer electrolyte additive"；"solid-state sulfur conversion"'),
        ('含液只写在方法里', '电池中另加液态电解液、正极液、离子液体界面层、凝胶；v2 只看标题摘要，分不出准固态'),
        ('"硼氢化物"', '还原剂 NaBH4、LiEt3BH，或合成前驱体'),
        ('"oxychloride"触发卤化物/固态', 'BiOCl 催化剂、锂/卤氧化物一次电池、氯离子电池'),
        ('"锂硫"误触发', '"LIBs""limits battery""Li2SO4""lithium/sulfur dioxide"；硫化聚丙烯腈只作前驱体（超级电容器、MoS2 生长）'),
        ('并列提及', '多体系综述里"锂硫"和"全固态"并列出现'),
    ]:
        S.append([a, b])
    sm = pd.DataFrame(S, columns=['项目', '结果'])

    note = pd.DataFrame({'说明': [
        '方案D 组别核查（2026-10-18），口径见 试点计划.md §D1。不改冻结的 v2 论文表。',
        '"组别核查_60篇"是复核表：一行一篇。v2 语义归组按分组本意判，并对照课题方案附录 D 冻结规则：违背本意且证据充分的记"不符合（已确认）"；附录 D 未覆盖或明文按标题判的固液混合等记"边界未明确"；多体系或以液态为主的综述、推断记"不符合（判断性）"。方案D 应用归属、主统计纳入、第二层相关性分开判。"规则层面原因"区分规则未按文字执行和规则字面命中（D8 已知局限），修规则时用。',
        '第二轮裁定（2026-10-18 晚）：#9 FeS2 正极单列为金属硫化物转化正极；#53 双归属；#18 待核实（残留溶剂需对应到电池所用样品）；#24 依据配方推定含液（保留 120 ℃ 真空处理步骤）。',
        '判据来源依次用：v2 摘要 → 库摘要 → 本地全文 → 只有标题。"依据原句"由代码到对应来源逐字核对（忽略空白、连字符、连字差异）。',
        '标记：✅ 原句找到且无需你定的事项 ｜🔶 原句找到但有需你定的事项（见"需人看的原因"）｜❌ 原句未找到。原句找到只说明文字在来源里，不说明判定正确。',
        '"定向排查标记清单"：关键词圈出候选，再逐篇看标题和摘要判定，是标记清单，不是已确认误收的下限。K4 只按标题粗分，只作线索。',
        '"敏感性情景清单"：固态锂硫组每篇被改的论文及其证据等级、情景 A/B 下的新组；随机样本的逐篇裁定优先于定向排查。脚本同时输出 敏感性情景清单.csv 供重算用。',
        '只报分类计数，不作错误率估计。60 篇只说明问题规模，不代表检索系统准确率，也查不到组外漏收。',
    ]})

    with pd.ExcelWriter(out, engine='openpyxl') as w:
        note.to_excel(w, '说明', index=False)
        sm.to_excel(w, '汇总', index=False)
        g.to_excel(w, '组别核查_60篇', index=False)
        k.to_excel(w, '定向排查标记清单', index=False)
        scdf.to_excel(w, '敏感性情景清单', index=False)
        widths = {'说明': {'A': 120}, '汇总': {'A': 52, 'B': 110},
                  '组别核查_60篇': {'E': 50, 'J': 22, 'K': 30, 'N': 30, 'P': 22, 'Q': 30, 'S': 24, 'W': 60, 'Z': 26, 'AA': 30, 'AB': 40},
                  '定向排查标记清单': {'A': 36, 'F': 70, 'G': 40}, '敏感性情景清单': {'B': 60, 'F': 24, 'G': 40, 'H': 14, 'I': 18, 'J': 18, 'K': 24}}
        for name, ws in w.sheets.items():
            for col in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col)].width = 14
            for col, wd in widths.get(name, {}).items():
                ws.column_dimensions[col].width = wd
            for row in ws.iter_rows():
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical='top')
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor='DDEBF7')
            ws.freeze_panes = 'B2' if name != '说明' else None
            if name in ('组别核查_60篇', '定向排查标记清单', '敏感性情景清单'):
                ws.auto_filter.ref = ws.dimensions
    print(sm.to_string(index=False))
    print('\n原句未找到：', g[g['原句核对'] == '未找到'][['序号', '依据原句']].to_string(index=False) if (g['原句核对'] == '未找到').any() else '无')


if __name__ == '__main__':
    main(*sys.argv[1:5])
