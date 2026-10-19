"""方案D 全文表样 v0.2：把 5 篇表样按表结构 v0.2（试点计划.md §D3，Z1–Z9）回填，并做证据复核。

用法：python build_fulltext_v02.py <全文文本目录> <输出.xlsx>
数据来源：build_fulltext_sample.py（v0.1）里 Claude 读全文填的数据 + 本文件里的 v0.2 改动和复核结果。
v0.2 改动：值与状态分开（主表只放值，例外记在"字段证据明细"）；行级证据注明支持字段；测试页加电池号、载量按电池号只填一处；
一行多样品/多窗口拆行；比较页数值列只放数，另列限定符、结果类型、值性质、取值方式；原文异常逐条编号；判定页 Z9 措辞；H01 降级。
"""
import importlib.util
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location('v01', HERE / 'build_fulltext_sample.py')
v01 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(v01)

GROUP = {p[0]: p[1] for p in v01.PAPERS}
DONE = '已抽取（正文 + SI）'
TOKENS = {'【未报告】': '查阅后未报告', '【待核查】': '待核查', '【不适用】': '不适用'}

# ======================= 样品页 =======================
S_FIELDS = ['文中样品名', '样品类型', '实验配置', '正极组成', '电解质', '负极', '施加的策略', '加工', '尺寸']
S_OVR = {  # 样品号: {字段: 新值}；'_file' '_page' '_fig' '_quote' 改行级证据
    'F1-S7': {'正极组成': 'VC/S : LPSCl1.5'},
    'F4-S3': {'正极组成': '同 F4-S2，VGCF 3.8 wt%'},
    'F4-S4': {'正极组成': 'VGCF : β-Li3PS4，VGCF 9.1 wt%（与复合正极中二者之比相同）', '负极': 'In 箔 8 mm（对电极兼参比）',
              '加工': '玛瑙研钵混合 15 min；30 mg β-Li3PS4 隔层 + 30 mg 复合物，30 kN（约 380 MPa）3 min', '实验配置': 'CV 电池（两电极）',
              '_file': 'SI', '_page': 8, '_fig': 'S-8 Cyclic Voltammetry', '_quote': 'corresponding to a VGCF/β-Li3PS4 composite with 9.1 mass-% VGCF'},
    'F5-S6': {'正极组成': '同 F5-S2，VGCF 先经 ALD 包覆 AlxOy（20 循环，2–3 nm）；包覆 VGCF 占 2.9 wt%'},
}
S_SUPPORT = {  # 行级原句支持的字段（未列出的默认"正极组成"）
    'F1-S4': '加工', 'F1-S5': '正极组成、加工', 'F1-S6': '正极组成', 'F2-S1': '加工', 'F2-S2': '正极组成、加工、施加的策略',
    'F2-S3': '正极组成、加工、施加的策略', 'F2-S4': '正极组成、加工、施加的策略', 'F3-S1': '正极组成', 'F3-S2': '文中样品名', 'F3-S3': '文中样品名',
    'F4-S4': '正极组成', 'F5-S6': '加工', 'F5-S7a': '正极组成',
}
S_SPLIT = {'F5-S7': [('F5-S7a', 'LPS/VGCF CV 复合'), ('F5-S7b', 'LPS/Super C65 CV 复合'), ('F5-S7c', 'LPS/CNT CV 复合'), ('F5-S7d', 'LPS/KBC CV 复合')]}

# ======================= 测试页 =======================
T_FIELDS = ['重复说明', '测试方法', '本电池载量', '运行压力', '温度', '电流程序', '电位窗口', '循环阶段/取值点']
BATT = {  # 测试号 → 电池号（= 同一配方、同一载量、同一装配方法的装配编号）
    'F1-S1-E1': 'F1-S1-B1', 'F1-S2-E1': 'F1-S2-B1', 'F1-S4-E1': 'F1-S4-B1',
    'F1-S2-E2': 'F1-S2-B2', 'F1-S3-E2': 'F1-S3-B2', 'F1-S4-E2': 'F1-S4-B2', 'F1-S2-E3': 'F1-S2-B2', 'F1-S4-E3': 'F1-S4-B2',
    'F1-S2-E4': 'F1-S2-B2', 'F1-S4-E4': 'F1-S4-B2', 'F1-S5-E1': 'F1-S5-B1', 'F1-S6-E1': 'F1-S6-B1', 'F1-S5-E2': 'F1-S5-B2',
    'F1-S6-E2': 'F1-S6-B2', 'F1-S5-E3': 'F1-S5-B3', 'F1-S5-E4': 'F1-S5-B4', 'F1-S6-E3': 'F1-S6-B3', 'F1-S7-E1': 'F1-S7-B1', 'F1-S7-E2': 'F1-S7-B1',
    'F2-S1-E1': 'F2-S1-B1', 'F2-S2-E1': 'F2-S2-B1', 'F2-S2-E3': 'F2-S2-B2', 'F2-S4-E1': 'F2-S4-B1', 'F2-S2-E4': 'F2-S2-B3',
    'F2-S3-E1': 'F2-S3-B1', 'F2-S4-E2': 'F2-S4-B2', 'F2-S2-E5': 'F2-S2-B2', 'F2-S5-E1': 'F2-S5-B1', 'F2-S2-E6': 'F2-S2-B1',
    'F3-S1-E1': 'F3-S1-B1', 'F3-S1-E2': 'F3-S1-B1', 'F3-S1-E3': 'F3-S1-B2', 'F3-S1-E4': 'F3-S1-B1', 'F3-S1-E5': 'F3-S1-B3',
    'F3-S1-E6': '不适用（吸附实验）', 'F3-S2-E1': 'F3-S2-B1', 'F3-S3-E1': 'F3-S3-B1',
    'F4-S1-E1': 'F4-S1-B1', 'F4-S2-E1': 'F4-S2-B1', 'F4-S3-E1': 'F4-S3-B1', 'F4-S1-E2': 'F4-S1-B1', 'F4-S2-E2': 'F4-S2-B1',
    'F4-S1-E3': 'F4-S1-B1', 'F4-S2-E3': 'F4-S2-B1', 'F4-S4-E1': 'F4-S4-B1',
    'F5-S2-E1': 'F5-S2-B1', 'F5-S3-E1': 'F5-S3-B1', 'F5-S1-E2': 'F5-S1-B1', 'F5-S2-E2': 'F5-S2-B1', 'F5-S6-E2': 'F5-S6-B1',
    'F5-S2-E3': '不适用（模拟）',
}
LOAD = {  # 电池号 → 载量（唯一来源；空 = 查阅后未报告）
    'F1-S1-B1': '1.1 mg cm⁻²（S）', 'F1-S2-B1': '1.3 mg cm⁻²（S）', 'F1-S4-B1': '1.6 mg cm⁻²（S）',
    'F1-S2-B2': '电解质约 4.5 mg cm⁻²', 'F1-S3-B2': '电解质约 4.5 mg cm⁻²', 'F1-S4-B2': '电解质约 4.5 mg cm⁻²',
    'F1-S5-B1': '1.5 mg cm⁻²', 'F1-S6-B1': '1.5 mg cm⁻²', 'F1-S5-B2': '1.7 mg cm⁻²', 'F1-S6-B2': '1.6 mg cm⁻²', 'F1-S5-B3': '3.4 mg cm⁻²',
    'F1-S5-B4': '7.6 mg cm⁻²；N/P 1.2', 'F1-S6-B3': '7.4 mg cm⁻²', 'F1-S7-B1': '',
    'F2-S1-B1': '—', 'F2-S2-B1': '15.3 mg cm⁻²（SE-C）', 'F2-S2-B2': '15.3 / 30.6 mg cm⁻²（厚度系列）', 'F2-S2-B3': '',
    'F2-S3-B1': '', 'F2-S4-B1': '', 'F2-S4-B2': '', 'F2-S5-B1': '', 'F2-S5-B2': '10.2 mg cm⁻²（S 3.40 mg cm⁻²）',
    'F3-S1-B1': '2.5 mg cm⁻²（原文"active material"）', 'F3-S1-B2': '4 mg cm⁻²（S）', 'F3-S1-B3': '', 'F3-S2-B1': '2.5 mg cm⁻²（S）', 'F3-S3-B1': '',
    'F4-S1-B1': '8.9 mg cm⁻²（NCM）', 'F4-S2-B1': '8.9 mg cm⁻²（原文通述；含碳后应略低，F5 同配方给 8.65）', 'F4-S3-B1': '', 'F4-S4-B1': '30 mg VGCF/SE 复合物（面载量未报告）',
    'F5-S1-B1': '8.92 mg cm⁻²', 'F5-S2-B1': '8.65 mg cm⁻²', 'F5-S3-B1': '8.65 mg cm⁻²', 'F5-S6-B1': '',
    'F5-S7a-B1': '', 'F5-S7b-B1': '', 'F5-S7c-B1': '', 'F5-S7d-B1': '',
    '不适用（吸附实验）': '不适用（吸附实验）', '不适用（模拟）': '不适用（模拟）',
}
F4I = '0.1C（原文：1.78 mA cm⁻²，见异常 A1）'
F5I = '0.1 C（正文：1.78 mA·cm⁻²；图 8 注：178 μA·m⁻²，见异常 A2、A3）'
TORQUE = '螺丝扭矩 10.7 N m'
T_OVR = {
    **{t: {'运行压力': TORQUE} for t in ['F1-S5-E1', 'F1-S6-E1', 'F1-S5-E2', 'F1-S6-E2', 'F1-S5-E3', 'F1-S5-E4', 'F1-S6-E3']},
    'F2-S1-E1': {'温度': '25 °C', '运行压力': '约 50 MPa'},
    'F3-S1-E3': {'电流程序': '正文 400 mA g⁻¹；SI 图 S7 注 500 mA g⁻¹（见异常 A4）'},
    'F4-S1-E1': {'电流程序': F4I}, 'F4-S2-E1': {'电流程序': F4I}, 'F4-S3-E1': {'电流程序': F4I},
    **{t: {'电位窗口': ''} for t in ['F2-S2-E3', 'F2-S4-E1', 'F2-S2-E5', 'F2-S5-E1']},
    'F2-S2-E6': {'电流程序': '0.1 mV s⁻¹ 扫到目标电位，保持 12 h', '电位窗口': '从 OCV（约 1.55 V vs In/InLi）到各目标电位（各点见 Fig. 6）'},
    'F4-S4-E1': {'测试方法': 'CV 后 XPS', '运行压力': '约 50 MPa', '温度': '', '电流程序': '1 mV s⁻¹ 扫描', '电位窗口': 'OCV → 5.5 V → 0 V vs In/InLi（两电极）',
                 '_file': 'SI', '_page': 8, '_fig': 'S-8 Cyclic Voltammetry', '_quote': 'the first sweep (oxidative sweep) was carried out to 5.5 V versus In/InLi'},
    **{t: {'电流程序': F5I} for t in ['F5-S2-E1', 'F5-S3-E1', 'F5-S1-E2', 'F5-S2-E2', 'F5-S6-E2']},
}
T_SUPPORT = {
    'F1-S1-E1': '运行压力、温度', 'F1-S2-E1': '运行压力、温度', 'F1-S4-E1': '运行压力、温度', 'F1-S2-E2': '本电池载量', 'F1-S3-E2': '本电池载量',
    'F1-S4-E2': '本电池载量', 'F1-S2-E4': '测试方法、重复说明', 'F1-S4-E4': '测试方法、重复说明', 'F1-S5-E1': '重复说明', 'F1-S6-E1': '重复说明',
    'F1-S6-E2': '本电池载量', 'F1-S5-E3': '循环阶段/取值点', 'F1-S5-E4': '本电池载量', 'F1-S6-E3': '本电池载量', 'F1-S7-E1': '电位窗口、温度',
    'F1-S7-E2': '电位窗口、温度', 'F2-S2-E1': '运行压力', 'F2-S2-E2a': '电流程序', 'F2-S5-E2a': '本电池载量', 'F3-S1-E2': '循环阶段/取值点',
    'F3-S1-E3': '电流程序、循环阶段/取值点', 'F3-S1-E4': '循环阶段/取值点', 'F3-S2-E1': '本电池载量', 'F4-S1-E1': '电流程序、本电池载量',
    'F2-S2-E2': '电流程序', 'F2-S5-E2': '本电池载量',
    'F4-S2-E1': '重复说明', 'F4-S1-E3': '重复说明、测试方法', 'F4-S2-E3': '重复说明、测试方法', 'F4-S4-E1': '电位窗口、电流程序',
    'F5-S2-E2': '循环阶段/取值点', 'F5-S6-E2': '循环阶段/取值点',
}
T_SPLIT = {
    'F1-S7-E1': [('F1-S7-E1', {'测试方法': '恒流长循环', '电位窗口': '上截止 2.7 V'}), ('F1-S7-E2', {'测试方法': '恒流长循环', '电位窗口': '上截止 3.1 V'})],
    'F2-S2-E2': [(f'F2-S2-E2{c}', {'测试方法': '恒流循环', '电位窗口': w + ' vs In/InLi'}) for c, w in zip('abcde', ['1–1.8 V', '1–2.4 V', '1–2.7 V', '1–3.0 V', '0–2.7 V'])],
    'F2-S5-E2': [(f'F2-S5-E2{c}', {'测试方法': '恒流循环', '电位窗口': w + ' vs In/InLi'}) for c, w in zip('abc', ['0–3.0 V', '1–3.0 V', '1–2.7 V'])],
    'F5-S7-E1': [(f'F5-S7{c}-E1', {}) for c in 'abcd'],
}
for c in 'abcde':
    BATT[f'F2-S2-E2{c}'] = 'F2-S2-B1'
for c in 'abc':
    BATT[f'F2-S5-E2{c}'] = 'F2-S5-B2'
for c in 'abcd':
    BATT[f'F5-S7{c}-E1'] = f'F5-S7{c}-B1'

# 论文级通用条件（一条原句覆盖同一论文某页某字段的多行；适用范围写在"适用于"）
PROV = [
    ('F1', '测试', ['运行压力', '温度'], 'F1-S1～S4 的测试（图 2 电池，40% 硫）', '正文', 'cycled at room temperature under a stack pressure of 50 MPa'),
    ('F1', '测试', ['电位窗口'], 'F1 全部恒流测试', '正文', 'Galvanostatic cycling was conducted in the voltage range of 1.4–3.1 V versus Li+/Li'),
    ('F1', '测试', ['运行压力'], 'F1-S5、S6 的测试（70% 硫电池，扭矩）', '正文', 'Screws on the casing were fastened by applying a torque of 10.7 N m'),
    ('F1', '测试', ['温度'], 'F1-S5、S6 室温测试（图 4）', '正文', 'SSSBs at room temperature. a,b, Charge–discharge voltage profiles of CNG–S'),
    ('F1', '样品', ['负极', '加工', '尺寸'], 'F1-S5、S6（70% 硫电池装配）', '正文', 'To assemble an SSSB, 70 mg of Li5.5PS4.5Cl1.5 was first pressed into a pellet in a 10-mm die'),
    ('F1', '样品', ['负极'], 'F1-S1～S4（图 2 电池）', '正文', 'Li–In|SE|VC–S–SE SSSBs at C/20'),
    ('F2', '测试', ['温度'], 'F2 全部循环测试', 'SI', 'All cycling tests were performed inside the climate chamber at 25 ℃'),
    ('F2', '测试', ['温度'], 'F2 全部阻抗测试', 'SI', 'all impedance measurements were performed in a climate chamber with the temperature controlled to be 25 ℃'),
    ('F2', '测试', ['运行压力'], 'F2 全部电池', 'SI', 'All cells were fixed in a metal frame to maintain a pressure at around 50 MPa'),
    ('F2', '样品', ['尺寸', '电解质'], 'F2-S2～S5', 'SI', 'loaded into the cell with a diameter of 10 mm (area of 0.785 cm2) as separator'),
    ('F3', '测试', ['电位窗口'], 'F3 全部恒流测试', '正文', 'in the voltage range of 1.5–2.8 V'),
    ('F3', '测试', ['运行压力'], 'F3 全部电池（扣式电池，不适用）', '正文', 'The CR2032 coin cells were fabricated'),
    ('F4', '测试', ['温度', '电位窗口'], 'F4 全部恒流测试', '正文', 'Cells with lithium metal anode are charged to 4.3 V and discharged to 2.6 V versus Li+/ Li at 25 °C'),
    ('F4', '测试', ['运行压力'], 'F4 全部电化学测试', '正文', 'applying a constant pressure of ∼50 MPa'),
    ('F4', '样品', ['加工'], 'F4-S1～S3', '正文', 'This powder stack was then compressed uniaxially at 30 kN (∼380 MPa) for 3 min'),
    ('F4', '测试', ['运行压力'], 'F4-S4 CV 电池', 'SI', 'A pressure of approximately 50 MPa was applied using an external frame around the cell casing'),
    ('F5', '测试', ['温度'], 'F5 全部电池测试', '正文', 'The battery tests were carried out at 25 °C'),
    ('F5', '测试', ['电位窗口'], 'F5 全部恒流测试', '正文', 'discharged to a lower cut-oﬀpotential of 2.6 V versus Li+/Li'),
    ('F5', '测试', ['运行压力'], 'F5 全部电化学测试', '正文', 'An external frame was used during all electrochemical tests, applying a constant pressure of approximately 50 MPa'),
    ('F5', '样品', ['加工'], 'F5-S1～S6', '正文', 'The resulting pellet was uniaxially compressed at 30 kN (approx. 380 MPa) for 3 min'),
    ('F5', '样品', ['电解质', '尺寸'], 'F5-S1～S6', '正文', 'LPS (60 mg) were added into a 10 mm diameter cylinder'),
    ('F5', '样品', ['负极'], 'F5 全部样品（锂箔或铟箔）', '正文', 'A 120 μm thick Li foil disc (Albemarle, Rockwood Lithium GmbH) or an indium foil disc'),
    ('F4', '样品', ['电解质'], 'F4-S1～S3', '正文', '60 mg of the β-Li3PS4 solid electrolyte was added'),
    ('F4', '样品', ['尺寸'], 'F4 全部样品', '正文', 'cylinder (inner diameter: 10 mm)'),
    ('F4', '样品', ['负极'], 'F4-S1～S3', '正文', 'A 120 μm thick lithium foil disc (Albemarle, Rockwood Lithium GmbH) with 9 mm diameter was used as anode'),
    ('F1', '样品', ['电解质'], 'F1-S1～S4（图 2 电池）', '正文', 'LPSCl1.5 or LPS as an electrolyte'),
    ('F3', '样品', ['负极', '电解质'], 'F3 全部样品', '正文', 'Li metal as a reference and counter electrode, and GF/D as a separator'),
    ('F3', '样品', ['尺寸', '加工'], 'F3-S1', '正文', '12 mm FRHS@CF electrodes are punched and calendered'),
]

# 人工字段明细（P1 复核结果、异常关联字段、定位与行级证据不同的字段）
# (记录号, 页, 字段, 值, 信息状态, 值性质, 取值方式, 文件, 原句或说明, 异常号)
MANUAL = [
    *[(s, '样品', '加工（装配压力）', '', '查阅后未报告', '', '', '正文', 'Methods 的装配步骤只写 70% 硫电池（F1-S5、S6）；40% 硫电池的装配压力未写', '') for s in ['F1-S1', 'F1-S2', 'F1-S3', 'F1-S4']],
    *[(t, '测试', '运行压力（MPa）', '', '查阅后未报告', '', '', '正文', '只给螺丝扭矩 10.7 N m，未换算成 MPa；第 2 页的 50 MPa 是对图 2 电池的陈述，是否适用于 70% 硫电池未写明', '')
      for t in ['F1-S5-E1', 'F1-S6-E1', 'F1-S5-E2', 'F1-S6-E2', 'F1-S5-E3', 'F1-S5-E4', 'F1-S6-E3']],
    ('F1-S5-E4', '测试', '温度', '60 °C', '已报告', '实验测量', '正文', '正文', 'At 60 °C, the', ''),
    ('F1-S6-E3', '测试', '温度', '60 °C', '待核查', '', '', 'SI', 'SI 图 28 的高载量 NG 电池温度：正文按 60 °C 对比，SI 图注未单独写', ''),
    ('F2-S1-E1', '测试', '运行压力', '约 50 MPa', '待核查', '', '', 'SI', '"All cells … around 50 MPa"写在 CV 段落；纯电解质片的阻抗测试是否同样加压未单独写', ''),
    ('F3-S1-E1', '测试', '本电池载量', '2.5 mg cm⁻²', '待核查', '', '', '正文', '本篇写"active material loading"，对照处写"loading of sulfur"；活性物是否指硫未写明', ''),
    ('F4-S3', '样品', '正极组成（比例基准）', '', '查阅后未报告', '', '', '正文', '3.8 wt% 的计算基准原文未重述（F4-S2 的 2.9 wt% 基准是正极总质量）', ''),
    ('F4-S4-E1', '测试', '温度', '', '查阅后未报告', '', '', 'SI', 'SI 的 CV 段落未写温度', ''),
    ('F5-S6', '样品', '正极组成（包覆 VGCF 质量分数）', '2.9 wt%', '已报告', '实验测量', '正文', '正文', 'uncoated, and with alumina-coated VGCF (2.9 wt %)', ''),
    ('F5-S6', '样品', '正极组成（包覆层质量是否计入）', '', '查阅后未报告', '', '', '正文', '原文只写包覆 VGCF 占 2.9 wt%，未说明包覆层质量是否计入', ''),
    *[(t, '测试', '电流程序', '0.1C', '冲突待核实', '', '正文', '正文', '0.1C 与原文给的电流密度 1.78 mA cm⁻² 不一致，见原文异常 A1', 'A1') for t in ['F4-S1-E1', 'F4-S2-E1', 'F4-S3-E1']],
    *[(t, '测试', '电流程序', '0.1 C', '冲突待核实', '', '正文', '正文', '正文 1.78 mA·cm⁻² 与图 8 注 178 μA·m⁻² 不一致，见原文异常 A2、A3', 'A2、A3') for t in ['F5-S2-E1', 'F5-S3-E1', 'F5-S1-E2', 'F5-S2-E2', 'F5-S6-E2']],
    ('F3-S1-E3', '测试', '电流程序', '', '冲突待核实', '', '正文；SI', '正文', '正文 400 mA g⁻¹，SI 图 S7 注 500 mA g⁻¹，见原文异常 A4', 'A4'),
    # ---- 开发案例涉及行的字段级定位（第二轮补）；已报告的第 9 项是原句，第 11 项是说明 ----
    ('F1-S4', '样品', '正极组成', 'CNG–S（40 wt% S）: LPSCl1.5 = 50:50', '已报告', '', '正文', '正文',
     'CNG–S, commercial Vulcan carbon–sulfur (VC–S) or NG–S (40 wt% S) were mixed with these two SEs (50:50 by weight)', ''),
    ('F1-S4', '样品', '正极组成（N 含量）', '15 at.%', '已报告', '实验测量', '正文', '正文', 'Nitrogen is evenly distributed within the CNG at a concentration of 15 at.%', ''),
    ('F2-S2', '样品', '负极', 'In 9 mm × 100 μm + Li 6 mm × 120 μm', '已报告', '', 'SI 正文', 'SI',
     'The employed In and Li foils possess the diameter of 9 mm and 6 mm and thickness of 100 μm and 120 μm', ''),
    *[(t, '测试', '电流程序', 'C/20', '已报告', '', '正文', '正文', 'to 4.1 V versus Li+/Li at C/20', '') for t in ['F1-S2-E2', 'F1-S3-E2', 'F1-S4-E2']],
    *[(t, '测试', '本电池载量', '电解质约 4.5 mg cm⁻²', '已报告', '', '正文', '正文', 'The LPSCl1.5 electrolyte loading was ~4.5 mg cm−2', '',
       '电池号 B2：TOF-SIMS 样品取自首次充电至 4.1 V 的电池（正文："were harvested from cells at an uncharged state (CNG OCV) and on initial charging to 4.1 V"）')
      for t in ['F1-S2-E4', 'F1-S4-E4']],
    ('F1-S5-E2', '测试', '重复说明', '有一个重复', '已报告', '', 'SI 图注', 'SI', 'e, Duplicate of cycling performance of the CNG/S cell at C/10', ''),
    ('F1-S5-E2', '测试', '本电池载量', '1.7 mg cm⁻²', '已报告', '', '图中图例', '正文', 'CNG–S, C/10, 1.7 mg cm–2', ''),
    ('F1-S5-E2', '测试', '电流程序', 'C/10', '已报告', '', '图中图例', '正文', 'CNG–S, C/10, 1.7 mg cm–2', ''),
    ('F1-S6-E2', '测试', '电流程序', 'C/10', '已报告', '', 'SI 图注', 'SI', 'Voltage profiles at C/10 with a sulfur loading of 1.6 mg cm-2', ''),
    *[(t, '测试', f, v, '已报告', '', 'SI 正文', 'SI', 'Cell(ii) and Cell(iii) were symmetric cells with either 15.3 or 30.6 mg cm-2 of SE-C composite on both sides of the cells', '')
      for t in ['F2-S2-E3', 'F2-S2-E5'] for f, v in [('重复说明', '两种厚度的对称电池'), ('本电池载量', '15.3 / 30.6 mg cm⁻²')]],
    *[(t, '测试', '电位窗口', '', '不适用', '', '', '', '阻抗在开路下测，不涉及电位窗口', '') for t in ['F2-S2-E3', 'F2-S4-E1', 'F2-S2-E5', 'F2-S5-E1']],
    ('F2-S2-E6', '测试', '重复说明', '还原、氧化各一个电池', '已报告', '', 'SI 正文', 'SI', 'the other one was used for the oxidation test', ''),
    ('F2-S2-E6', '测试', '本电池载量', '15.3 mg cm⁻²', '已报告', '', 'SI 正文', 'SI', 'followed by loading 12 mg (15.3 mg cm-2) of SE-C composite', ''),
    *[('F2-S2-E6', '测试', f, '', '已报告', '', 'SI 正文', 'SI',
       'scanned from OCV (about 1.55 V vs. In/InLi) to the target potential with 0.1 mV s-1 and kept at the target potential for 12 h', '') for f in ['电流程序', '电位窗口']],
    ('F4-S2-E1', '测试', '本电池载量', '8.9 mg cm⁻²', '已报告', '', '正文', '正文', 'for cells with an active material loading of 8.9 mg cm−2', '',
     '原文对全部电池通述；含碳后 NCM 载量应略低（F5 同配方给 8.65），未单独报告'),
    ('F4-S4-E1', '测试', '本电池载量', '30 mg 复合物', '已报告', '', 'SI 正文', 'SI', '30 mg of the VGCF/SE composite was added as working electrode', ''),
    ('F5-S1-E2', '测试', '本电池载量', '8.92 mg cm⁻²', '已报告', '', '正文', '正文', 'cathode composite corresponds to an AM loading of 8.92 mg·cm−2', ''),
    ('F5-S2-E2', '测试', '本电池载量', '8.65 mg cm⁻²', '已报告', '', '正文', '正文', 'mg of composite refer to 8.65 mg·cm−2', ''),
]

# ======================= 比较页 =======================
# 比较号: (样品测试号, 对照测试号, 对照来源, 结果类型, 样品值限定, 样品值, 对照值限定, 对照值, 单位, 定性描述, 误差或不确定性, 值性质, 取值方式, 受异常影响, 复核说明)
C_V2 = {
    'F1-C1': (None, None, '同篇', '定量', '', 1210, '', 1130, 'mAh g⁻¹', '', '', '实验测量', '正文', '', ''),
    'F1-C2': (None, None, '同篇', '定性', '', None, '', None, '', '样品：充电约 2.9 V 处无额外平台；对照：有', '', '实验测量', '图（定性）', '', '限制属实：样品与对照倍率不同（C/10 vs C/20）'),
    'F1-C3': (None, None, '同篇', '定量', '约', 46, '', 277, 'mAh g⁻¹（按电解质质量）', '', '', '实验测量', '正文', '', ''),
    'F1-C4': (None, None, '同篇', '定量', '', 200, '', 277, 'mAh g⁻¹（按电解质质量）', '', '', '实验测量', '正文', '', ''),
    'F1-C5': (None, None, '同篇', '定性', '', None, '', None, '', '样品：只剩一个宽而弱的峰；对照：约 2.4 V、2.94 V 两步氧化', '', '实验测量', '图（定性）', '', ''),
    'F1-C6': (None, None, '同篇', '定性', '', None, '', None, '', '样品：P2S7⁻ 未检出；对照：明显', '基质效应可能使强度略有偏差（作者说明）', '实验测量', '图（定性）', '', ''),
    'F1-C7': (None, None, '同篇', '定量', '', 3.67e5, '', 1.27e6, 'Ω', '', '', '实验测量（等效电路拟合）', 'SI 正文', '', ''),
    'F1-C8': (None, None, '同篇', '定量（仅样品值）', '', 480, '', None, 'mAh g⁻¹', '原文：较高倍率下 CNG 容量约为 NG 的 2.5 倍，未指明具体倍率，不能据此推算 1C 时 NG 的值', '3 个电池，平均差 < 10%', '实验测量', '正文', '',
              '已复核：对照值只在图中，不影响候选判断，不读图'),
    'F1-C9': (None, None, '同篇', '定量', '', 990, '', 480, 'mAh g⁻¹', '', '', '实验测量', '正文', '', ''),
    'F1-C10': (None, None, '同篇', '定量', '', 11.3, '', 4.1, 'mAh cm⁻²', '', '', '实验测量', '正文', '', ''),
    'F1-C11': ('F1-S7-E1', 'F1-S7-E2', '同篇', '定性', '', None, '', None, '', '上截止 2.7 V：长循环仍很差；3.1 V：差', '', '实验测量', 'SI 图（定性）', '', '限制属实：数值只在图中；S7 硫含量未报告'),
    'F2-C1': (None, None, '同篇', '定量', '', 0.20, '', 0.76, 'mS cm⁻¹', '', '样品 ±0.02', '实验测量', '正文', '', ''),
    'F2-C2': (None, None, '同篇', '定量', '', 0.033, '', 0.76, 'mS cm⁻¹', '', '', '实验测量', '正文', '', ''),
    'F2-C3a': ('F2-S3-E1', 'F2-S2-E4', '同篇', '定量', '', 160, '', 327, 'mAh gSE⁻¹', '', '', '实验测量', '正文', '', ''),
    'F2-C3b': ('F2-S4-E2', 'F2-S2-E4', '同篇', '定量', '', 87, '', 327, 'mAh gSE⁻¹', '', '', '实验测量', '正文', '', ''),
    'F2-C4': (None, None, '同篇', '定量', '', 0.52, '', 0.41, 'eV', '对照两个对称电池分别为 0.41、0.43 eV', '作者称在典型标准差内比较', '实验测量', '正文', '', ''),
    'F2-C5': ('F2-S2-E6', 'F2-S2-E6', '同篇（同一电池逐点）', '定量', '>', 1500, '', 40, 'Ω', '', '', '实验测量', '正文', '', ''),
    'F2-C6': ('F2-S2-E2d', 'F2-S2-E2a', '同篇', '定量（仅样品值）', '约', 200, '', None, 'mAh gSE⁻¹', '对照：很少（原文未给数值）', '', '实验测量', '正文', '', ''),
    'F2-C7a': ('F2-S5-E2c', 'F2-S5-E2a', '同篇', '定量（仅样品值）', '', 106.9, '', None, '%', '样品 1350 mAh gS⁻¹、4.59 mAh cm⁻²；对照（0–3.0 V）：快速衰减', '另一批材料复现（Fig. S14）', '实验测量', '正文', '',
               '已复核：对照值只在图中，不影响候选判断，不读图'),
    'F2-C7b': ('F2-S5-E2c', 'F2-S5-E2b', '同篇', '定量（仅样品值）', '', 106.9, '', None, '%', '对照（1–3.0 V）：后期库仑效率低', '同上', '实验测量', '正文', '', '已复核：同上'),
    'F3-C1': (None, None, '同篇', '定量', '', 740, '', 504, 'mAh g⁻¹', '', '', '实验测量', '正文', '', '限制属实：对照组成未报告'),
    'F3-C2': (None, None, '同篇', '定量', '', 73, '', 65, '%', '取值点不同：样品第 325 圈（以第 25 圈为基准），对照第 130 圈（基准未写）', '', '实验测量', '正文', '', '限制属实：取值点不可比'),
    'F3-C3': (None, None, '作者引用的历史结果', '定量', '', 1165, '', 1100, 'mAh g⁻¹', '', '', '实验测量', 'SI 图注', '', '限制属实：跨论文对照，条件未核对'),
    'F3-C4': (None, None, '同篇（同一电池前后）', '定量', '', 8.23, '', 3.4, 'Ω cm²', '', '', '实验测量（拟合）', '表格', '', ''),
    'F3-C5': (None, None, '同篇（同一电池前后）', '定量', '', 870, '', 1154, 'mAh g⁻¹', '', '', '实验测量', '正文', '', '限制属实：无不含 FeS2 的对照'),
    'F4-C1': (None, None, '同篇', '定量', '', 144, '', 98, 'mAh g⁻¹', '', '重复范围：样品 144–155，对照 98–128', '实验测量', '正文', 'A1（两侧电流相同，组内比较不受影响）', ''),
    'F4-C2': (None, None, '同篇', '定性', '', None, '', None, '', '样品第 62 圈后容量低于参比', '', '实验测量', '正文', 'A1（两侧电流相同）', ''),
    'F4-C3': (None, None, '同篇', '定量', '', 98.9, '', 99.5, '%', '', '', '实验测量', '正文', 'A1（两侧电流相同）', ''),
    'F4-C4': (None, None, '同篇', '定量', '', 227, '', 201, 'mAh g⁻¹', '', '', '实验测量', '正文', 'A1（两侧电流相同）', ''),
    'F4-C5': (None, None, '同篇', '定量（仅样品值）', '约', 60, '', None, '%', '对照占比只在图中', '', '实验测量', '正文', '', '已复核：对照值只在图中，不影响候选判断，不读图'),
    'F4-C6': (None, None, '同篇', '定性', '', None, '', None, '', '样品：多硫化物碎片增加更明显；对照：增加', '每样 5 个点', '实验测量', '图（定性）', '', ''),
    'F4-C7': (None, None, '同篇', '定性', '', None, '', None, '', 'CV 后出现与复合正极相似的分解信号', '', '实验测量', 'SI 图（定性）', '', ''),
    'F5-C1a': ('F5-S7b-E1', 'F5-S7a-E1', '同篇', '定性', '', None, '', None, '', 'C65（62 m² g⁻¹）分解信号高于 VGCF（24 m² g⁻¹）', '', '实验测量', '图（定性）', '', ''),
    'F5-C1b': ('F5-S7d-E1', 'F5-S7a-E1', '同篇', '定性', '', None, '', None, '', 'KBC（1270 m² g⁻¹）分解信号高于 VGCF', '', '实验测量', '图（定性）', '', ''),
    'F5-C1c': ('F5-S7c-E1', 'F5-S7a-E1', '同篇', '定性', '', None, '', None, '', '作者称 CNT 是"随比表面积增大"规律的例外', 'CNT 电池峰电流波动大（作者说明）', '实验测量', '图（定性）', '', ''),
    'F5-C2': (None, None, '同篇', '定量（差值）', '约', 30, '', None, 'mAh g⁻¹（VGCF 比 C65 高出的值）', '第 18 圈后两者相同', '', '实验测量', '正文', 'A2、A3（两侧电流相同）', ''),
    'F5-C3': (None, None, '同篇', '定量', '', 29, '约', 46, '%', '', '对照为读图值：首圈约 122、第 50 圈约 56 mAh g⁻¹，估读约 ±3 mAh g⁻¹（约 ±3 个百分点）',
              '样品：实验测量；对照：由读图值换算', '样品：正文；对照：读图（图 8a）', 'A2、A3（两侧电流相同）', '已复核：读图取对照值'),
    'F5-C4': (None, None, '同篇', '定量', '', 52, '', 29, '%', '', '', '实验测量', '正文', 'A2、A3（两侧电流相同）', ''),
    'F5-C5': (None, None, '同篇', '定量', '', 1.4, '', 1, '倍（第 50 圈放电容量，相对无碳）', '', '', '实验测量', '正文', 'A2、A3（两侧电流相同）', ''),
    'F5-C6a': ('F5-S2-E2', 'F5-S1-E2', '同篇', '定量', '约', 650, '约', 560, 'Ω', '', '', '实验测量（拟合）', '正文', '', ''),
    'F5-C6b': ('F5-S2-E2', 'F5-S6-E2', '同篇', '定量', '约', 650, '约', 560, 'Ω', '', '', '实验测量（拟合）', '正文', '', ''),
    'F5-C7a': ('F5-S2-E2', 'F5-S1-E2', '同篇', '定量', '约', 93, '约', 74, 'Ω', '', '', '实验测量（拟合）', '正文', '', ''),
    'F5-C7b': ('F5-S2-E2', 'F5-S6-E2', '同篇', '定量', '约', 93, '约', 74, 'Ω', '', '', '实验测量（拟合）', '正文', '', ''),
    'F5-C8': (None, '无碳虚拟电极', '同篇', '定量（增幅）', '约', 27, '', None, '%（容量增幅）', '', '', '理论计算或模拟', '正文', '', ''),
}
C_SPLIT = {'F2-C3': ['F2-C3a', 'F2-C3b'], 'F2-C7': ['F2-C7a', 'F2-C7b'], 'F5-C1': ['F5-C1a', 'F5-C1b', 'F5-C1c'], 'F5-C6': ['F5-C6a', 'F5-C6b'], 'F5-C7': ['F5-C7a', 'F5-C7b']}
C_COND = {  # 测试条件核对改写
    'F2-C1': '温度 一致（25 °C，SI 写明阻抗统一在 25 ℃ 测）；压力 无法判断（纯电解质片是否加压未单独写）；其余 一致',
    'F2-C2': '同 F2-C1',
    'F2-C3a': '一致（SI 图 S9 注：三种配比同一方式循环，1–3 V vs In/InLi）', 'F2-C3b': '同 F2-C3a',
    'F1-C9': '倍率 一致（C/10）；载量 明确不同但接近（1.7 vs 1.6）；温度 一致；装配 一致（同一扭矩）',
    'F5-C1a': '一致', 'F5-C1b': '一致', 'F5-C1c': '一致',
}
REF_FIX = [('F2-C7', 'F2-C7a、C7b'), ('F5-C4～C7', 'F5-C4、C5、C6a、C6b、C7a、C7b'), ('F5-C1、F5-C2', 'F5-C1a–c、F5-C2'), ('F2-C3', 'F2-C3a、C3b')]

# ======================= 原文异常 =======================
ANOM = [
    ('A1', 'F4', '电流程序（F4 全部恒流测试）', '正文 p2（Experimental）；正文 p5（Results，Fig. 1B 前）——一条异常、两个定位',
     '0.1C, corresponding to 1.78 mA cm−2 for cells with an active material', '正文',
     'C 率与电流密度不一致', '0.178 mA cm⁻²', '0.1 × 8.9 mg cm⁻² × 200 mAh g⁻¹ = 0.178 mA cm⁻²',
     '载量 8.9 mg cm⁻²：F4 正文 p2；理论容量 200 mAh g⁻¹：F4 正文 p5（"assumption: Qtheo = 200 mAh g−1"）',
     '冲突待核实', 'F4 各恒流测试的电流程序；F4-C1～C4 两侧电流相同，组内比较不受影响；涉及电流数值的跨论文比较暂缓'),
    ('A2', 'F5', '电流程序（F5 全部恒流测试）', '正文 p3（Experimental）',
     'which corresponds to 1.78 mA·cm−2 for cells', '正文',
     '同一句里 0.1 C、1.78 mA·cm⁻² 与 1.78 mA h·cm⁻² 不自洽', '0.178 mA cm⁻²', '0.1 × 1.78 mAh cm⁻² = 0.178 mA cm⁻²',
     '面容量 1.78 mA h·cm⁻²：同一句；理论容量 200 mA h·g⁻¹：F5 正文 p2', '冲突待核实（与 A3 不一致）',
     'F5 各恒流测试的电流程序；F5-C2～C5 两侧电流相同，组内比较不受影响。F4 与 F5 为同一课题组、措辞相同，可能沿用同一模板（仅推测）'),
    ('A3', 'F5', '电流程序（图 8 电池）', '正文 p10（图 8 图注）',
     'the current density was 178 μA·m−2', '正文',
     '单位写作 μA·m⁻²', '0.178 mA cm⁻²（= 178 μA cm⁻²）', '同 A2',
     '同 A2', '疑似单位笔误；改成 μA·cm⁻² 后与复算一致（未确认）', '同 A2'),
    ('A4', 'F3', '电流程序（4 mg cm⁻² 高载量测试）', '正文 p4（Fig. 6a 段落）；SI 图 S7 图注',
     'till 120 cycles at 400 mA g−1', '正文',
     '正文 400 mA g⁻¹，SI 图 S7 注 500 mA g⁻¹', '—（无法复算判断哪个对）', '—',
     '正文与 SI 各一处', '冲突待核实', 'F3-S1-E3 的电流程序；没有比较行用到该测试'),
]
ANOM_SI_QUOTE = {'A4': ('SI', 'The areal specific capacity (mAh cm-2) plot of high sulfur loading of 4 mg cm-2 at 500 mA g-1')}

# ======================= 判定页 / 组合假设 =======================
J_FIX = {
    'J01': {'固态锂硫应用情况': '在规定检索范围内检出：F1、F2（范围：本次 5 篇全文；2026-10-18）'},
    'J02': {'固态锂硫应用情况': '尚未检索（本次只看 5 篇；正式做时按检索式查固态锂硫全库，写"在规定检索范围内检出 / 未检出"，记检索日期、关键词、范围）'},
    'J03': {'固态锂硫应用情况': '尚未检索（摘要层约 88 篇提到催化或介体，未逐篇核对是否用 FeS2）'},
    'J04': {'固态锂硫应用情况': '尚未检索'},
    'J05': {'固态锂硫应用情况': '在规定检索范围内检出：F1（范围：本次 5 篇全文；2026-10-18）'},
}
H01 = {
    '能否独立操控': '原则上可，待配方约束确认',
    '待定配方约束': '①碳含量变化时替代哪一组分：固定 S:SE 比例外加碳，还是碳替代部分 SE 或 S（F4、F5 为外加）；②包覆层质量是否计入碳的质量分数（F5 只写包覆 VGCF 占 2.9 wt%，包覆层是否计入未报告）；③硫面载量如何保持一致（建议固定硫面载量，正极总质量随碳含量变）',
    '指标 Y（方向、尺度）': '主指标：第 50 圈放电容量（原值，越大越好）。辅助：①电解质氧化容量（越小越好），按 F1 的独立测试配置测——单独装含硫 40% 的宿主–硫–电解质电池，Li–In 对电极，C/20 直接充电至 4.1 V vs Li⁺/Li（按 In/InLi = 0.62 V vs Li⁺/Li 约合 3.48 V vs In/InLi），按电解质质量计容量，不接在主循环测试之后；②首圈放电容量；③TLM 有效离子/电子电导',
    '优先级': '候选之一，不享有预设优先级；主问题和主实验在批2 与候选比较后定',
}


# ======================= 设计决策开发案例涉及的行（J01 → J02 → H01：导电碳在固态锂硫硫正极中的取舍）=======================
CASE_CMP = ['F1-C3', 'F1-C6', 'F1-C9', 'F2-C1', 'F2-C2', 'F2-C5', 'F4-C2', 'F4-C7', 'F5-C3', 'F5-C4', 'F5-C5', 'F5-C6a', 'F5-C6b', 'F5-C7a', 'F5-C7b']
CASE_TEST = ['F1-S4-E2', 'F1-S2-E2', 'F1-S4-E4', 'F1-S2-E4', 'F1-S5-E2', 'F1-S6-E2', 'F2-S2-E3', 'F2-S1-E1', 'F2-S4-E1', 'F2-S2-E6',
             'F4-S2-E1', 'F4-S1-E1', 'F4-S4-E1', 'F5-S2-E2', 'F5-S1-E2', 'F5-S6-E2']
CASE_SAMPLE = sorted({t.rsplit('-', 1)[0] for t in CASE_TEST})
CASE_ROWS = set(CASE_CMP + CASE_TEST + CASE_SAMPLE)
CASE_ROLE = {
    'F1-C3': '固态锂硫内：限制界面电子传输的宿主（CNG）使电解质氧化容量从 277 降到约 46', 'F1-C6': '同上，分解产物 P2S7⁻ 碎片未检出',
    'F1-C9': '同上，第 50 圈容量 990 vs 480（NG 对照）', 'F2-C1': '代价：加 10 wt% 碳，有效离子电导降到 0.20（无碳 0.76）',
    'F2-C2': '代价：30 wt% 碳，有效离子电导降到 0.033', 'F2-C5': '代价：2.4 V vs In/InLi 保持 12 h 后，复合层离子电阻 40 → >1500 Ω',
    'F4-C2': '代价（非硫固态）：加 VGCF 的初始优势在第 62 圈后消失', 'F4-C7': '代价（非硫固态）：碳/电解质界面本身在电位扫描下分解',
    'F5-C3': '代价（非硫固态）：未包覆 VGCF 第 50 圈保持 29%，低于无碳约 46%（读图）', 'F5-C4': '改造（非硫固态）：包覆后保持 52% vs 未包覆 29%',
    'F5-C5': '改造：包覆 VGCF 第 50 圈容量为无碳的 1.4 倍', 'F5-C6a': '代价：未包覆使 R_SE 约 560 → 650 Ω', 'F5-C6b': '改造：包覆后 R_SE 回到约 560 Ω',
    'F5-C7a': '代价：未包覆使 R_Ca 约 74 → 93 Ω', 'F5-C7b': '改造：包覆后 R_Ca 回到约 74 Ω',
}
CASE_STEPS = [
    ('1 决策问题', '固态锂硫的硫/硫化物复合正极必须加导电碳提供电子通路，但碳/硫化物电解质界面会加速电解质分解。能否借用非硫固态中"绝缘包覆导电碳纤维"的做法，在保留电子通路的同时减少分解？',
     '方案D §D0 第二层范围：正极 + 正极—电解质界面；负极只记条件', ''),
    ('2 目标配置', '硫/Li6PS5Cl（或 LPSCl1.5）/碳复合正极，In/InLi 或 Li–In 负极，室温，约 50 MPa，约 1.4–3.1 V vs Li⁺/Li', '判定页 J01、J02', ''),
    ('3 观察：碳的代价', '固态锂硫中碳降低有效离子电导、在高电位下使离子电阻激增；非硫固态中碳的初始优势随循环消失，碳/电解质界面本身分解，未包覆碳使循环保持率低于无碳',
     'F2-C1、C2、C5；F4-C2、C7；F5-C3、C6a、C7a', '见"开发案例_证据"页'),
    ('4 观察：改造的效果（非硫固态）', 'ALD AlxOy 包覆 VGCF 后，第 50 圈保持率 52%（未包覆 29%），容量为无碳的 1.4 倍，R_SE、R_Ca 回到无碳水平', 'F5-C4、C5、C6b、C7b', ''),
    ('5 观察：固态锂硫内的类比', '用宽带隙 C3N4 限制宿主/电解质界面的局部电子传输，电解质氧化减少、容量提高——支持"限制界面电子传输而不切断整体电子通路"可行',
     'F1-C3、C6、C9', '类比级别：功能相近、材料不同（宿主改性 vs 包覆碳）'),
    ('6 作者解释（策略页）', '①碳/电解质界面的电位降是分解的主要驱动力；②包覆要薄，纤维端头和压制破裂处仍保持电子接触；③只适用纤维状碳。①②③均为作者推测或假设',
     'F5-ST2；F4-ST1、F5-ST1、F2-ST1、F1-ST1', '前提的支持证据为"推测"，不是实验验证'),
    ('7 新假设（判定页）', 'J01：碳在固态锂硫中已有应用，代价在固态锂硫中同样被观察到。J02：包覆碳纤维进硫正极 = 改造候选，适用性未知——前提②（端头接触够不够硫的大面积转化）和③（包覆层在硫正极工作区间是否稳定）未知',
     'J01、J02', '"已有应用"的检索范围只有本次 5 篇（Z9）'),
    ('8 待检验的问题', '包覆后电解质分解是否减少？同时硫利用率是否因碳表面电子供应受阻而下降？', 'J02', ''),
    ('9 可区分的实验（候选）', '碳含量 × 有无包覆四组（H01）。开做前须先定配方约束：碳含量变化替代哪一组分、包覆层质量是否计入、硫面载量如何固定；电解质氧化容量按独立测试配置测（直接充至 4.1 V vs Li⁺/Li）',
     'H01', 'H01 不享有预设优先级'),
    ('10 证据状态', '案例涉及的样品、测试数据字段均已有定位或状态。保留意见：F5-C3 对照为读图值；F4、F5 电流数值有冲突（A1–A3，组内比较两侧相同，不受影响，但不用于跨论文电流比较）；'
     'F2-S1-E1 的压力待核查；F5 包覆层质量是否计入未报告；F1 70% 硫电池的压力未报告（只有扭矩）', '原文异常页；字段证据明细页', ''),
    ('11 本案例不代表', '只用 5 篇演示"观察 → 作者解释 → 新假设 → 可区分实验"的方法链；不是研究结论，不决定第二层主问题或主实验', '', ''),
]


# ======================= 工具 =======================
def norm(s):
    s = unicodedata.normalize('NFKC', str(s))
    return re.sub(r'[\s\-‐‑‒–—―−­]', '', s).lower()


class Texts:
    """规范化后的全文 + 每段规范化文本对应的页码（SI 由 docx 转来时没有页码）。"""
    def __init__(self, d):
        self.n = {}
        for p in Path(d).glob('*.txt'):
            buf, spans, pos, cur = [], [], 0, None
            for part in re.split(r'(=== 第 \d+ 页 ===)', p.read_text(encoding='utf-8')):
                m = re.fullmatch(r'=== 第 (\d+) 页 ===', part)
                if m:
                    cur = int(m.group(1))
                    continue
                s = norm(part)
                buf.append(s)
                spans.append((pos, pos + len(s), cur))
                pos += len(s)
            self.n[p.stem] = (''.join(buf), spans)

    def find(self, paper, file, quote):
        if not quote:
            return '无原句', None
        key = paper if file == '正文' else f'{paper}_SI'
        if key not in self.n:
            return '原文未找到', None
        s, pages = self.n[key]
        i = s.find(norm(quote))
        if i < 0:
            return '原文未找到', None
        pg = next((p for a, b, p in pages if a <= i < b), None)
        return '程序核验通过', pg


def strip_tokens(v):
    """返回 (值, 信息状态或 None, 原写法)。"""
    if v is None:
        return '', None, ''
    raw = str(v)
    if raw.strip() in ('—', '-', ''):
        return '', ('不适用' if raw.strip() in ('—', '-') else None), raw
    status = None
    for tok, st in TOKENS.items():
        if tok in raw:
            status = st
    rd = re.search(r'【读图[^】]*】', raw)
    if rd:
        status = '只在图中'
    clean = re.sub(r'【[^】]*】', '', raw).strip(' ，；、')
    if re.fullmatch(r'（[^）]*）', clean or ''):
        clean = ''
    return clean, status, raw if status else ''


def main(text_dir, out):
    T = Texts(text_dir)
    detail = []

    def add_detail(rid, page, field, value, status, nature='', how='', file='', note='', anom='', src='自动', quote=''):
        paper = rid[:2]
        ver, pg = T.find(paper, file, quote) if quote else ('无原句', None)
        detail.append({'记录号': rid, '页': page, '字段': field, '值': value, '信息状态': status, '值性质': nature, '取值方式': how,
                       '文件': file, '页码': pg, '原句或说明': quote or note, '说明': note if quote else '', '异常号': anom,
                       '核验状态': ver if quote else ('不需核验' if status != '已报告' else ('读图，无原句' if '读图' in how else '见行级原句')),
                       '来源': src, '复核 ✍': ''})

    # ---------- 样品 ----------
    srows = []
    for r in v01.SAMPLES:
        sid = r[0]
        base = dict(zip(['样品号'] + S_FIELDS + ['_summary', '_file', '_page', '_fig', '_quote'], r))
        base.update(S_OVR.get(sid, {}))
        targets = S_SPLIT.get(sid, [(sid, None)])
        for nid, name in targets:
            row = dict(base)
            row['样品号'] = nid
            if name:
                row['文中样品名'] = name
                row['正极组成'] = 'LPS : 碳，碳 9.1 wt%（与全电池中二者之比相同）'
            srows.append(row)
    S = []
    for row in srows:
        sid = row['样品号']
        rec = {'样品号': sid, '论文': sid[:2], '组': GROUP[sid[:2]]}
        for f in S_FIELDS:
            val, st, raw = strip_tokens(row[f])
            rec[f] = val
            if st and not any(m[0] == sid and m[2].startswith(f) for m in MANUAL):
                add_detail(sid, '样品', f, val, st, note=f'v0.1 写法：{raw}')
        ver, pg = T.find(sid[:2], row['_file'], row['_quote'])
        rec.update({'抽取状态': DONE, '文件': row['_file'], '页码': pg or row['_page'], '图表': row['_fig'], '原句': row['_quote'],
                    '支持字段': S_SUPPORT.get(sid, '正极组成'), '核验状态': ver})
        S.append(rec)
    S = pd.DataFrame(S)

    # ---------- 测试 ----------
    trows = []
    for r in v01.TESTS:
        tid = r[0]
        base = dict(zip(['测试号', '样品号', '重复说明', '测试方法', '本电池载量', '运行压力', '温度', '电流程序', '电位窗口', '循环阶段/取值点',
                         '_summary', '_file', '_page', '_fig', '_quote'], r))
        base.update({k: v for k, v in T_OVR.get(tid, {}).items()})
        for nid, ov in T_SPLIT.get(tid, [(tid, {})]):
            row = dict(base)
            row['测试号'] = nid
            if nid.startswith('F5-S7'):
                row['样品号'] = nid[:6]
            row.update(ov)
            trows.append(row)
    Tt = []
    for row in trows:
        tid = row['测试号']
        bid = BATT[tid]
        rec = {'测试号': tid, '样品号': row['样品号'], '论文': tid[:2], '电池号': bid}
        for f in T_FIELDS:
            if f == '本电池载量':
                val, st = LOAD[bid], (None if LOAD[bid] else '查阅后未报告')
                if bid.startswith('不适用'):
                    st = None
                if val == '—':
                    val, st = '', '不适用'
                raw = ''
            else:
                val, st, raw = strip_tokens(row[f])
            rec[f] = val
            manual_hit = any(m[0] == tid and m[2].startswith(f.split('/')[0]) for m in MANUAL)
            if st and not manual_hit:
                add_detail(tid, '测试', f, val, st, note=(f'v0.1 写法：{raw}' if raw else ('电池号 ' + bid + ' 的载量未报告' if f == '本电池载量' else '')))
        ver, pg = T.find(tid[:2], row['_file'], row['_quote'])
        rec.update({'抽取状态': DONE, '文件': row['_file'], '页码': pg or row['_page'], '图表': row['_fig'], '原句': row['_quote'],
                    '支持字段': T_SUPPORT.get(tid, T_SUPPORT.get(re.sub(r'[a-e]$', '', tid), '测试方法')), '核验状态': ver})
        Tt.append(rec)
    Tt = pd.DataFrame(Tt)
    load_chk = Tt.groupby('电池号')['本电池载量'].nunique()
    load_bad = load_chk[load_chk > 1]

    for m in MANUAL:
        rid, page, field, value, status, nature, how, file, note, anom, *extra = m
        quote = note if status == '已报告' else ''
        add_detail(rid, page, field, value, status, nature, how, file, (extra[0] if extra else '') if quote else note, anom, '人工', quote)

    # ---------- 比较 ----------
    C = []
    for r in v01.COMPARISONS:
        cid = r[0]
        b = dict(zip(['比较号', '样品测试号', '对照测试号', '对照来源', '实际变化', '单因素', '测试条件核对', '指标', '样品值', '对照值', '单位', '取值点',
                      '重复数', '误差或不确定性', '观察方向', '目标意义', '值来源', '需人看的原因', '_file', '_page', '_fig', '_quote'], r))
        for nid in C_SPLIT.get(cid, [cid]):
            v = C_V2[nid]
            rec = {'比较号': nid, '论文': nid[:2], '样品测试号': v[0] or b['样品测试号'], '对照测试号': v[1] or b['对照测试号'], '对照来源': v[2],
                   '实际变化': b['实际变化'], '单因素': b['单因素'], '测试条件核对': C_COND.get(nid, b['测试条件核对']), '指标': b['指标'],
                   '结果类型': v[3], '样品值限定': v[4], '样品值': v[5], '对照值限定': v[6], '对照值': v[7], '单位': v[8], '定性描述': v[9],
                   '取值点': b['取值点'], '重复数': strip_tokens(b['重复数'])[0], '误差或不确定性': v[10] or strip_tokens(b['误差或不确定性'])[0],
                   '观察方向': b['观察方向'], '目标意义': b['目标意义'], '值性质': v[11], '取值方式': v[12], '受异常影响': v[13]}
            if nid == 'F2-C5':
                rec['取值点'] = '2.4 V vs In/InLi（样品）vs 初始（对照）'
            if nid in ('F5-C1a', 'F5-C1b', 'F5-C1c'):
                rec['实际变化'] = '碳种类 VGCF → ' + {'F5-C1a': 'Super C65', 'F5-C1b': 'KBC', 'F5-C1c': 'CNT'}[nid] + '（比表面积与形貌同时变）'
                rec['观察方向'] = '无法判断' if nid == 'F5-C1c' else '增加'
            if nid in ('F5-C6b', 'F5-C7b'):
                rec['实际变化'] = '未包覆 VGCF vs 包覆 VGCF（同为 2.9 wt%）'
            if nid in ('F5-C6a', 'F5-C7a'):
                rec['实际变化'] = '外加未包覆 VGCF 2.9 wt% vs 无碳'
            if nid == 'F2-C3b':
                rec['实际变化'] = '碳 10 → 30 wt%'
            if nid == 'F2-C3a':
                rec['实际变化'] = '碳 10 → 20 wt%'
            for f in ('重复数', '误差或不确定性'):
                if '【未报告】' in str(b[f]) and not rec[f]:
                    add_detail(nid, '比较', f, '', '查阅后未报告')
            if rec['结果类型'].startswith('定量') and rec['对照值'] is None and '仅样品值' in rec['结果类型']:
                add_detail(nid, '比较', '对照值', '', '只在图中，未读图' if '图' in (v[9] + v[14]) else '查阅后未报告', note=v[14] or v[9])
            if nid == 'F5-C3':
                add_detail(nid, '比较', '对照值', '约 46 %', '已报告', '由读图值换算', '读图（图 8a）', '正文',
                           '读图方式：按 y 轴主刻度（50 mAh g⁻¹）目测"No carbon additive"曲线首圈约 122、第 50 圈约 56 mAh g⁻¹，46% = 56/122；估读约 ±3 mAh g⁻¹。'
                           '自洽检查：同图未包覆 VGCF 约 44/150 = 29%、包覆约 80/155 = 52%、80/56 = 1.4，均与正文一致', src='人工')
            if v[11] == '理论计算或模拟':
                add_detail(nid, '比较', '样品值', v[5], '已报告', '理论计算或模拟', '正文', '正文', '微结构分辨的电化学模拟（虚拟电极），不是实测', src='人工')
            ver, pg = T.find(nid[:2], b['_file'], b['_quote'])
            rec.update({'抽取状态': DONE, '文件': b['_file'], '页码': pg or b['_page'], '图表': b['_fig'], '原句': b['_quote'],
                        '支持字段': '样品值、对照值' if rec['结果类型'] == '定量' else ('样品值' if rec['结果类型'].startswith('定量') else '定性描述'),
                        '核验状态': ver, '复核说明': v[14]})
            C.append(rec)
    C = pd.DataFrame(C)

    # ---------- 策略 ----------
    stg_cols = ['策略号', '策略家族', '具体实现', '位置', '针对的问题', '声称作用', '支持作用的比较', '支持层次', '作用所需前提', '前提提出者', '前提的支持证据',
                '代价或限制', '代价证据', '跨体系相关陈述', '文件', '页码', '原句']
    G = pd.DataFrame(v01.STRATEGIES, columns=stg_cols)
    for col in ('支持作用的比较', '前提的支持证据', '代价证据'):
        for a, b2 in REF_FIX:
            G[col] = G[col].str.replace(a, b2, regex=False)
    G.insert(1, '论文', G['策略号'].str[:2])
    G['抽取状态'] = DONE
    sup = {'F1-ST1': '跨体系相关陈述', 'F1-ST2': '具体实现、作用所需前提', 'F2-ST1': '针对的问题', 'F2-ST2': '具体实现、声称作用', 'F2-ST3': '具体实现',
           'F2-ST4': '代价或限制', 'F3-ST1': '具体实现', 'F3-ST2': '作用所需前提', 'F4-ST1': '跨体系相关陈述', 'F5-ST1': '作用所需前提', 'F5-ST2': '作用所需前提'}
    G['支持字段'] = G['策略号'].map(sup)
    res = G.apply(lambda r: pd.Series(T.find(r['论文'], r['文件'], r['原句'])), axis=1)
    G['核验状态'], G['页码'] = res[0], res[1].fillna(G['页码'])

    # ---------- 判定 / 组合假设 / 异常 ----------
    J = pd.DataFrame(v01.JUDGMENTS, columns=['判定号', '策略号', '来源组', '目标配置', '比较级别', '类比依据与差异', '前提逐项核对', '固态锂硫应用情况', '稳定性说明',
                                             '代价证据', '判定', '待检验的问题', '依据（比较号）'])
    for jid, fix in J_FIX.items():
        for k, v in fix.items():
            J.loc[J['判定号'] == jid, k] = v
    for col in ('依据（比较号）', '前提逐项核对', '待检验的问题'):
        for a, b2 in REF_FIX:
            J[col] = J[col].str.replace(a, b2, regex=False)
    J.loc[J['判定号'] == 'J02', '待检验的问题'] = J.loc[J['判定号'] == 'J02', '待检验的问题'].str.replace(
        '直接充电至 4.1 V 的电解质氧化容量', '电解质氧化容量（独立测试配置，直接充电至 4.1 V vs Li⁺/Li，见 H01）', regex=False)
    H = pd.DataFrame(v01.COMBOS, columns=['假设号', '策略甲', '策略乙', '目标配置', '能否独立操控', '假设内容', '文献依据', '实验组', '指标 Y（方向、尺度）', '重复数', '需要的机制表征'])
    for k, v in H01.items():
        H.loc[H['假设号'] == 'H01', k] = v
    H.loc[H['假设号'] == 'H01', '目标配置'] = '硫/Li6PS5Cl/VGCF 复合正极，In/InLi 负极，25 °C，约 50 MPa，1–2.7 V vs In/InLi（约 1.62–3.32 V vs Li⁺/Li）'
    for a, b2 in REF_FIX:
        H['文献依据'] = H['文献依据'].str.replace(a, b2, regex=False)

    A = []
    for a in ANOM:
        aid, paper, field, locs, quote, file, issue, recalc, formula, inputs, status, impact = a
        ver, pg = T.find(paper, file, quote)
        row = {'异常号': aid, '论文': paper, '字段': field, '出现位置': locs, '原文': quote, '核验状态': ver, '页码（第一处）': pg, '疑点': issue,
               '复算值': recalc, '公式': formula, '输入来源': inputs, '状态': status, '影响': impact, '复核 ✍': ''}
        if aid in ANOM_SI_QUOTE:
            f2, q2 = ANOM_SI_QUOTE[aid]
            row['原文'] += f'｜SI：{q2}'
            row['核验状态'] += '；SI ' + T.find(paper, f2, q2)[0]
        A.append(row)
    A = pd.DataFrame(A)

    P = []
    for paper, page, fields, scope, file, quote in PROV:
        ver, pg = T.find(paper, file, quote)
        P.append({'论文': paper, '页': page, '字段': '、'.join(fields), '适用于': scope, '文件': file, '页码': pg, '原句': quote, '核验状态': ver})
    P = pd.DataFrame(P)

    D = pd.DataFrame(detail)

    # ---------- 结构验证 ----------
    cov_batt = {r['电池号'] for _, r in Tt.iterrows() if '本电池载量' in str(r['支持字段'])
                or ((D['记录号'] == r['测试号']) & D['字段'].str.startswith('本电池载量') & (D['信息状态'] == '已报告')).any()}
    batt_ok = lambda r, f, pg: pg == '测试' and f == '本电池载量' and r.get('电池号') in cov_batt

    def coverage(df, idcol, fields, page):
        cnt = {'行级原句': 0, '字段明细': 0, '论文级通用条件': 0, '电池级定位': 0, '引用其他行': 0, '空白且有状态': 0, '空白无状态': 0, '有值无定位': 0}
        miss = []
        for _, r in df.iterrows():
            rid = r[idcol]
            for f in fields:
                v = r[f]
                has_d = ((D['记录号'] == rid) & D['字段'].str.startswith(f.split('/')[0])).any() if len(D) else False
                has_p = any(p[0] == rid[:2] and p[1] == page and f in p[2] for p in PROV)
                if v in ('', None) or (isinstance(v, float) and pd.isna(v)):
                    cnt['空白且有状态' if has_d else '空白无状态'] += 1
                    if not has_d:
                        miss.append(f'{rid}:{f}（空白无状态）')
                elif f in str(r['支持字段']):
                    cnt['行级原句'] += 1
                elif has_d:
                    cnt['字段明细'] += 1
                elif str(v).startswith('同'):
                    cnt['引用其他行'] += 1
                elif has_p:
                    cnt['论文级通用条件'] += 1
                elif batt_ok(r, f, page):
                    cnt['电池级定位'] += 1
                else:
                    cnt['有值无定位'] += 1
        return cnt, miss
    S_DATA = ['正极组成', '电解质', '负极', '加工', '尺寸']              # 要求可追溯到原文的数据字段
    T_DATA = ['重复说明', '本电池载量', '运行压力', '温度', '电流程序', '电位窗口']
    cs, ms = coverage(S, '样品号', S_DATA, '样品')
    ct, mt = coverage(Tt, '测试号', T_DATA, '测试')
    nodef_s = coverage(S, '样品号', S_DATA, '样品')[0]['有值无定位']
    gap = []
    for df, idc, flds, pg in [(S, '样品号', S_DATA, '样品'), (Tt, '测试号', T_DATA, '测试')]:
        for _, r in df.iterrows():
            for f in flds:
                v = r[f]
                if v in ('', None) or str(v).startswith('同') or f in str(r['支持字段']):
                    continue
                if ((D['记录号'] == r[idc]) & D['字段'].str.startswith(f.split('/')[0])).any():
                    continue
                if any(p[0] == r[idc][:2] and p[1] == pg and f in p[2] for p in PROV):
                    continue
                if batt_ok(r, f, pg):
                    continue
                gap.append((r[idc][:2], pg, f))
    gap_rows = []
    for df, idc, flds, pg in [(S, '样品号', S_DATA, '样品'), (Tt, '测试号', T_DATA, '测试')]:
        for _, r in df.iterrows():
            for f in flds:
                v = r[f]
                if v in ('', None) or str(v).startswith('同') or f in str(r['支持字段']):
                    continue
                if ((D['记录号'] == r[idc]) & D['字段'].str.startswith(f.split('/')[0])).any():
                    continue
                if any(p[0] == r[idc][:2] and p[1] == pg and f in p[2] for p in PROV):
                    continue
                if batt_ok(r, f, pg):
                    continue
                gap_rows.append({'记录号': r[idc], '页': pg, '字段': f, '值': v, '属开发案例': '是' if r[idc] in CASE_ROWS else '否',
                                 '处理': '待补定位；补完前不标证据完整，不用于确定性比较'})
    GAP = pd.DataFrame(gap_rows, columns=['记录号', '页', '字段', '值', '属开发案例', '处理'])
    for df, idc in [(S, '样品号'), (Tt, '测试号')]:
        cnt = GAP['记录号'].value_counts()
        df['证据完整性'] = df[idc].map(lambda x: f'未完整（缺 {cnt[x]} 个定位）' if x in cnt else '数据字段均有定位或状态')
    gap = pd.Series(gap).value_counts() if gap else pd.Series(dtype=int)
    num_bad = C[C['结果类型'].str.startswith('定量') & ~C['样品值'].apply(lambda x: isinstance(x, (int, float)))]['比较号'].tolist()
    allv = pd.concat([S[['核验状态']], Tt[['核验状态']], C[['核验状态']], G[['核验状态']]])
    V = [
        ('行数', f'样品 {len(S)}｜测试 {len(Tt)}｜比较 {len(C)}｜策略 {len(G)}｜判定 {len(J)}｜组合假设 {len(H)}｜原文异常 {len(A)}｜字段证据明细 {len(D)}｜论文级通用条件 {len(P)}'),
        ('拆行（Z8 及同类）', 'F5-S7 → S7a–d；F5-S7-E1 → 4 行；F1-S7-E1 → E1/E2；F2-S2-E2 → E2a–e；F2-S5-E2 → E2a–c；F2-C3、F2-C7、F5-C6、F5-C7 → a/b；F5-C1 → a–c'),
        ('行级原句核验（只说明原句能在来源中找到）', '；'.join(f'{k} {v}' for k, v in allv['核验状态'].value_counts().items())),
        ('论文级通用条件原句核验', '；'.join(f'{k} {v}' for k, v in P['核验状态'].value_counts().items())),
        ('字段明细中带原句的核验', '；'.join(f'{k} {v}' for k, v in D['核验状态'].value_counts().items())),
        ('原文异常原句核验', '；'.join(A['核验状态'])),
        ('电池号载量一致性（同一电池号只能有一个载量）', '通过' if load_bad.empty else '不一致：' + '、'.join(load_bad.index)),
        ('比较页数值列只放数（结果类型为定量的行）', '通过' if not num_bad else '不合规：' + '、'.join(num_bad)),
        ('样品页数据字段定位覆盖（正极组成、电解质、负极、加工、尺寸）', '；'.join(f'{k} {v}' for k, v in cs.items())),
        ('测试页数据字段定位覆盖（重复、载量、压力、温度、电流、窗口）', '；'.join(f'{k} {v}' for k, v in ct.items())),
        ('"有值无定位"集中在哪（论文·页·字段：个数）', '；'.join(f'{a}·{b}·{c}：{n}' for (a, b, c), n in gap.items()) or '无'),
        ('定位缺口合计（样品页 + 测试页）', f"{len(GAP)} 个；其中属开发案例 {(GAP['属开发案例'] == '是').sum()} 个（明细见“定位缺口”页）"),
        ('不要求原句的字段', '样品名、样品类型、实验配置、施加的策略、测试方法、循环阶段/取值点：分析者归类或描述，定位随行级证据'),
        ('空白无状态的字段', '；'.join(ms + mt) or '无'),
        ('字段明细的信息状态分布', '；'.join(f'{k} {v}' for k, v in D['信息状态'].value_counts().items())),
        ('比较页结果类型', '；'.join(f'{k} {v}' for k, v in C['结果类型'].value_counts().items())),
        ('比较页值性质', '；'.join(f'{k} {v}' for k, v in C['值性质'].value_counts().items())),
        ('比较页对照来源', '；'.join(f'{k} {v}' for k, v in C['对照来源'].value_counts().items())),
    ]
    V = pd.DataFrame(V, columns=['检查项', '结果'])

    # ---------- 开发案例 ----------
    CS = pd.DataFrame(CASE_STEPS, columns=['步骤', '内容', '依据', '保留意见'])
    comp = {}
    for df, idc in [(S, '样品号'), (Tt, '测试号')]:
        comp.update(dict(zip(df[idc], df['证据完整性'])))
    ce = C[C['比较号'].isin(CASE_CMP)].copy()
    ce['组'] = ce['论文'].map(GROUP)
    ce['在案例中的作用'] = ce['比较号'].map(CASE_ROLE)

    def rows_ok(r):
        ids = [x.strip() for x in re.split(r'[/、]', f"{r['样品测试号']}/{r['对照测试号']}") if x.strip()]
        ids += [i.rsplit('-', 1)[0] for i in ids]
        bad = [i for i in ids if i in comp and comp[i] != '数据字段均有定位或状态']
        return '样品、测试数据字段均有定位或状态' if not bad else '未完整：' + '、'.join(bad)
    ce['样品与测试证据'] = ce.apply(rows_ok, axis=1)
    ce['保留意见'] = ce.apply(lambda r: '；'.join(x for x in [r['取值方式'] if '读图' in r['取值方式'] else '', f"受异常影响 {r['受异常影响']}" if r['受异常影响'] else '',
                                                               '' if r['单因素'] in ('是',) else f"单因素：{r['单因素']}"] if x), axis=1)
    CE = ce[['比较号', '组', '在案例中的作用', '指标', '样品值限定', '样品值', '对照值限定', '对照值', '单位', '定性描述', '观察方向', '目标意义', '值性质', '样品与测试证据', '保留意见']]

    # ---------- 证据复核 ----------
    R = pd.DataFrame(REVIEW, columns=['v0.1 行号', '优先级', '复核结论', '复核状态'])

    note = pd.DataFrame({'说明': [
        '方案D 全文表样 v0.2（5 篇，2026-10-18）：按表结构 v0.2（试点计划.md §D3）回填 v0.1 的数据，并按优先级做证据复核。v0.1 文件保留不动。',
        '主表只放值。例外（查阅后未报告、待核查、不适用、只在图中、读图、换算、模拟、冲突待核实，以及定位与行级原句不同的字段）记在"字段证据明细"。空白不等于未报告：空白字段要么在明细里有状态，要么在"结构验证"里列为"空白无状态"。',
        '行级原句只支持"支持字段"列出的字段；方法部分的通用条件（温度、压力、电压窗口、装配）按论文记在"论文级通用条件"，一条原句覆盖适用范围内的多行。',
        '测试页的电池号 = 同一配方、同一载量、同一装配方法的装配编号；载量按电池号只填一处，代码检查是否一致。',
        '比较页：样品值、对照值只放数；"约""＞"放在限定列；定性结果写在"定性描述"。值性质分实验测量 / 由实测值换算 / 理论计算或模拟；取值方式分正文 / 表格 / 图注 / 读图。',
        '原文异常按条编号（A1–A4），出现位置另记；不替作者改；冲突未解决前，相关条件不支持确定性的定量比较。',
        '"结构验证"报表结构是否装得下；"证据复核"报 v0.1 的 45 行需人看（44 🔶 + 1 🧮）按优先级复核到哪一步。两者分开：结构验证通过不等于证据复核完成。',
        '原句核验通过只说明原句能在 PDF 文本里找到，不说明组别、标签、数值归属、比较条件或机制解释正确。',
        '第二轮（2026-10-18 晚）：定位缺口按优先级补——先补设计决策开发案例涉及的行（已补齐）；载量按电池号定位；阻抗测试的电位窗口记不适用。其余缺口列在“定位缺口”页，补完前对应行不标证据完整，不用于确定性比较。',
        '设计决策开发案例按方案D §8.2a 写在 设计决策案例_D01.md，不扩展工作簿；“定位缺口”页的“属开发案例”列标出案例涉及的行。',
    ]})

    with pd.ExcelWriter(out, engine='openpyxl') as w:
        for name, df in [('说明', note), ('结构验证', V), ('证据复核', R), ('样品', S), ('测试', Tt), ('比较', C), ('策略', G), ('判定', J), ('组合假设', H),
                         ('原文异常', A), ('字段证据明细', D), ('论文级通用条件', P), ('定位缺口', GAP)]:
            if name not in ('说明', '结构验证') and '复核 ✍' not in df.columns:
                df = df.assign(**{'复核 ✍': ''})
            df.to_excel(w, name, index=False)
        for ws in w.book.worksheets:
            ws.freeze_panes = 'B2'
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:80])
                c0 = col[0].column
                ws.column_dimensions[c0 if isinstance(c0, str) else get_column_letter(c0)].width = min(max(8, width * 1.1), 46)
            for row in ws.iter_rows():
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical='top')
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor='DDEBF7')
        w.book['说明'].column_dimensions['A'].width = 130
        w.book['结构验证'].column_dimensions['B'].width = 120
    print(V.to_string(index=False))
    print('\n证据复核：', R['复核状态'].value_counts().to_dict())
    print('优先级 × 状态：'); print(pd.crosstab(R['优先级'], R['复核状态']).to_string())
    bad = pd.concat([S[S['核验状态'] != '程序核验通过'][['样品号', '原句']].rename(columns={'样品号': '号'}),
                     Tt[Tt['核验状态'] != '程序核验通过'][['测试号', '原句']].rename(columns={'测试号': '号'}),
                     C[C['核验状态'] != '程序核验通过'][['比较号', '原句']].rename(columns={'比较号': '号'}),
                     G[G['核验状态'] != '程序核验通过'][['策略号', '原句']].rename(columns={'策略号': '号'}),
                     P[P['核验状态'] != '程序核验通过'][['论文', '原句']].rename(columns={'论文': '号'}),
                     D[D['核验状态'] == '原文未找到'][['记录号', '原句或说明']].rename(columns={'记录号': '号', '原句或说明': '原句'})])
    print('\n原句未找到：'); print(bad.to_string() if len(bad) else '无')


# ======================= 证据复核（v0.1 的 45 行需人看） =======================
# 优先级：P1 影响分组、关键数值、核心比较、H01/J01/J02 依据；P2 其余比较和策略依据；P3 其余
REVIEW = [
    ('F5-C3', 'P1', '对照值读图取得约 46%（首圈约 122、第 50 圈约 56 mAh g⁻¹）；同图另两条曲线读数与正文 29%、52%、1.4 倍一致', '已复核：已解决'),
    ('F1-S1', 'P1', '运行压力 50 MPa 原文明示（p2，图 2 电池）；装配压力查阅后未报告（Methods 只写 70% 硫电池）', '已复核：已解决'),
    ('F1-S2', 'P1', '同 F1-S1（F1-C1、C3 的对照样品）', '已复核：已解决'),
    ('F1-S3', 'P1', '同 F1-S1（F1-C4 的样品）', '已复核：已解决'),
    ('F1-S5', 'P1', '装配 2 t/3 t/1 t、10 mm 模具、扭矩 10.7 N m 原文明示（p10）；MPa 未报告', '已复核：已解决'),
    ('F1-S6', 'P1', '同 F1-S5（F1-C9 的对照）', '已复核：已解决'),
    ('F1-S5-E2', 'P1', '运行压力改为"螺丝扭矩 10.7 N m"；p2 的 50 MPa 针对图 2 电池，不套用；F1-C9 两侧装配相同', '已复核：已解决'),
    ('F1-S6-E2', 'P1', '同 F1-S5-E2', '已复核：已解决'),
    ('F2-S1-E1', 'P1', '温度 25 °C（SI：阻抗统一在 25 ℃ 测）；压力仍待核查（"All cells … 50 MPa"写在 CV 段落）', '已复核：部分解决（压力待你看）'),
    ('F4-S4', 'P1', '从 SI p8 补齐：VGCF 9.1 wt%、玛瑙研钵 15 min、30 mg 隔层 + 30 mg 复合物、380 MPa 3 min、In 箔 8 mm 对电极', '已复核：已解决'),
    ('F4-S4-E1', 'P1', '从 SI p8 补齐：OCV→5.5 V→0 V vs In/InLi，1 mV s⁻¹，两电极，约 50 MPa；温度查阅后未报告', '已复核：已解决'),
    ('F5-S6', 'P1', '包覆 VGCF 占 2.9 wt% 原文明示（p10）；包覆层质量是否计入查阅后未报告（影响 H01 配方约束②）', '已复核：已解决'),
    ('F4-S1-E1', 'P1', '电流数值与 C 率不一致，记异常 A1，电流程序记"冲突待核实"', '已复核：已解决'),
    ('F4-S2-E1', 'P1', '同 F4-S1-E1', '已复核：已解决'),
    ('F5-S2-E1', 'P1', '记异常 A2、A3，电流程序记"冲突待核实"', '已复核：已解决'),
    ('F1-C2', 'P2', '倍率不同（C/10 vs C/20）属实，单因素"否"保留，比较只作定性参考', '已复核：限制属实'),
    ('F1-C8', 'P2', '"2.5 倍"原文指较高倍率泛称，未指明 1C，不能推算 NG 的 1C 值；对照值只在图中，不影响候选判断，不读图', '已复核：限制属实'),
    ('F1-C11', 'P2', '数值只在图中，S7 硫含量未报告，属实', '已复核：限制属实'),
    ('F2-C3', 'P2', '条件核对补齐（SI 图 S9：同一方式循环，1–3 V vs In/InLi）；拆成 C3a、C3b', '已复核：已解决'),
    ('F2-C7', 'P2', '对照值只在图中，不影响候选判断，不读图；拆成 C7a、C7b', '已复核：限制属实'),
    ('F4-C5', 'P2', '对照占比只在图中，不影响候选判断（J01 未引用），不读图', '已复核：限制属实'),
    ('F4-S3', 'P2', '3.8 wt% 的计算基准原文未重述，记"查阅后未报告"', '已复核：已解决'),
    ('F5-S7', 'P2', '拆成 S7a–d（VGCF、C65、CNT、KBC）', '已复核：已解决'),
    ('F5-S7-E1', 'P2', '拆成 4 个测试', '已复核：已解决'),
    ('F5-S2-E3', 'P2', '模拟（虚拟电极），值性质记"理论计算或模拟"', '已复核：已解决'),
    ('F5-C8', 'P2', '同 F5-S2-E3', '已复核：已解决'),
    ('F1-S5-E1', 'P2', '运行压力同 F1-S5-E2 处理', '已复核：已解决'),
    ('F1-S6-E1', 'P2', '同上', '已复核：已解决'),
    ('F1-S5-E3', 'P2', '同上', '已复核：已解决'),
    ('F1-S5-E4', 'P2', '同上；60 °C 原文明示', '已复核：已解决'),
    ('F1-S6-E3', 'P2', '运行压力同上；SI 图 28 的温度未单独写，待核查', '已复核：部分解决（温度待你看）'),
    ('F1-S7', 'P3', '硫含量与加工查阅后未报告', '未复核（P3）'),
    ('F1-S7-E1', 'P3', '已拆成 E1/E2；其余未复核', '未复核（P3）'),
    ('F3-S1', 'P3', '比例写法相加为 90，待你看', '未复核（P3）'),
    ('F3-S2', 'P3', '对照组成未报告', '未复核（P3）'),
    ('F3-S3', 'P3', '跨论文对照', '未复核（P3）'),
    ('F3-S1-E1', 'P3', '"活性物"是否指硫：原文两处说法不同，记待核查', '已复核：部分解决（待你看）'),
    ('F3-S1-E3', 'P3', '记异常 A4，电流程序"冲突待核实"', '已复核：已解决'),
    ('F3-S1-E5', 'P3', '无对照电池', '未复核（P3）'),
    ('F3-S1-E6', 'P3', '无空白对照', '未复核（P3）'),
    ('F3-S3-E1', 'P3', '跨论文', '未复核（P3）'),
    ('F3-C1', 'P3', '对照组成未报告，属实', '已复核：限制属实'),
    ('F3-C2', 'P3', '取值点不可比，属实', '已复核：限制属实'),
    ('F3-C3', 'P3', '对照来源改为"作者引用的历史结果"', '已复核：限制属实'),
    ('F3-C5', 'P3', '无对照电池，属实', '已复核：限制属实'),
]

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
